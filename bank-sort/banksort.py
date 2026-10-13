# -*- coding: utf-8 -*-
"""은행정리: 수동으로 받은 신한은행 조회·이체결과, 홈택스 매출·매입 세금계산서 엑셀을
은행 폴더의 통합 파일에 더하고 「은행내역 통합」을 다시 만든다.

엑셀의 [은행정리] 버튼(은행정리.xlsm) → AppleScriptTask(banksort.scpt) → 이 파일.
    python banksort.py 파일1 파일2 ...      (파일 없이 실행하면 통합만 다시 만든다)
결과 문구는 표준출력으로 돌려준다. 첫 줄이 OPEN:<경로> 이면 버튼이 그 파일을 연다.
"""
import datetime as dt
import json
import os
import re
import runpy
import shutil
import sys
import traceback
import unicodedata
import warnings
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")
HERE = Path(__file__).resolve().parent
HOME = Path(os.environ.get("BANKSORT_HOME", Path.home() / "Library/Application Support/BankSort"))
CONFIG = HOME / "config.json"

OUT_NAME = "은행내역 통합 2024-2026.xlsx"
BUY_NAME = "매입세금계산서 통합 2024-2026.xlsx"
SELL_NAME = "매출세금계산서 통합 2024-2026.xlsx"
Q_EXTRA = "조회 추가분.xlsx"
T_EXTRA = "이체결과 추가분.xlsx"
REPORT_SHEET = "이번 정리"

N = lambda s: unicodedata.normalize("NFC", str(s))


def load_config():
    try:
        return json.loads(CONFIG.read_text("utf-8"))
    except Exception:
        return {}


def save_config(cfg):
    HOME.mkdir(parents=True, exist_ok=True)
    CONFIG.write_text(json.dumps(cfg, ensure_ascii=False, indent=1), "utf-8")


def base_dir():
    env = os.environ.get("BANKSORT_BASE")
    if env:
        return Path(env)
    return Path(load_config().get("base") or Path.home() / "Desktop/은행")


# ------------------------------------------------------------------ 엑셀 읽기


def read_rows(path):
    """시트별 행 목록(값만). .xlsx 와 예전 .xls 모두."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        try:
            book = xlrd.open_workbook(str(path))
        except Exception:
            return read_html_table(path)
        out = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_EMPTY or v == "":
                        v = None
                    row.append(v)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    # read_only 는 시트에 적힌 크기를 믿는데, 은행 내려받기 파일은 크기가 1×1 로 적혀 있어 쓰지 않는다
    wb = openpyxl.load_workbook(str(path), data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def read_html_table(path):
    """확장자만 .xls 인 HTML 표(일부 은행·홈택스 내려받기)."""
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__(); self.rows = []; self.row = None; self.cell = None
        def handle_starttag(self, tag, attrs):
            if tag == "tr": self.row = []
            elif tag in ("td", "th") and self.row is not None: self.cell = []
        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.cell is not None:
                t = " ".join("".join(self.cell).split()); self.row.append(t or None); self.cell = None
            elif tag == "tr" and self.row is not None:
                self.rows.append(self.row); self.row = None
        def handle_data(self, data):
            if self.cell is not None: self.cell.append(data)

    raw = path.read_bytes()
    for enc in ("utf-8", "cp949", "euc-kr"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError:
            continue
    p = P(); p.feed(text)
    return [("sheet1", p.rows)]


def norm_header(v):
    return re.sub(r"\s+", "", N(v)) if v is not None else ""


def find_table(rows, required):
    """required 머리글이 모두 있는 첫 행을 찾아 (머리글 목록, 아래 행들)."""
    for i, row in enumerate(rows[:30]):
        hdr = [norm_header(v) for v in row]
        if all(r in hdr for r in required):
            return hdr, rows[i + 1:]
    return None, None


def num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    s = str(v).replace(",", "").replace("원", "").strip()
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return 0


def num_or_none(v):
    """홈택스는 수량·단가를 '7,000,000' 같은 글자로 준다."""
    if v in (None, ""):
        return None
    return num(v)


def clean(v):
    return v.strip() if isinstance(v, str) else v


def to_ts(v):
    if isinstance(v, dt.datetime):
        return v
    s = str(v).strip()
    for fmt in ("%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y.%m.%d %H:%M", "%Y-%m-%d %H:%M"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"거래일시를 읽지 못했습니다: {v!r}")


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = re.sub(r"[./]", "-", str(v).strip())[:10]
    return dt.datetime.strptime(s, "%Y-%m-%d").date()


# ------------------------------------------------------------------ 신한은행 원본

Q_REQ = ["적요", "입금액", "출금액", "내용", "잔액"]
T_REQ = ["이체구분", "처리결과", "수취인성명", "처리금액"]
Q_COLS = ["No", "전체선택", "거래일시", "적요", "입금액", "출금액", "내용", "잔액", "거래점명", "입금인코드", "메모"]
T_COLS = ["No", "전체선택", "이체구분", "거래일시", "처리결과", "출금계좌", "입금은행", "입금계좌", "수취인성명", "처리금액",
          "입금금액", "수수료", "오류금액", "등록일자", "오류코드", "거래메모", "입금인코드", "출금통장표시내용", "입금통장표시내용",
          "수취인수수료부담여부"]


def table_dicts(rows, required):
    hdr, body = find_table(rows, required)
    if hdr is None:
        return None
    out = []
    for r in body:
        d = {h: (r[i] if i < len(r) else None) for i, h in enumerate(hdr) if h}
        out.append(d)
    return out


def bank_q_rows(rows):
    """조회 원본 → 표준 행(dict). 거래일시가 둘로 나뉜 양식도 받는다."""
    hdr, _ = find_table(rows, Q_REQ)
    if hdr is None:
        return None
    if "거래일시" not in hdr and not ("거래일자" in hdr and ("거래시간" in hdr or "시간" in hdr)):
        return None
    out = []
    for d in table_dicts(rows, Q_REQ):
        when = d.get("거래일시")
        if when is None and d.get("거래일자"):
            when = f'{N(d.get("거래일자")).strip()} {N(d.get("거래시간") or d.get("시간") or "00:00:00").strip()}'
        if when is None or (d.get("입금액") is None and d.get("출금액") is None):
            continue
        try:
            t = to_ts(when)
        except ValueError:
            continue
        out.append(dict(t=t, 적요=clean(d.get("적요")), 입금=num(d.get("입금액")), 출금=num(d.get("출금액")),
                        내용=clean(d.get("내용")) or "", 잔액=num(d.get("잔액")), 점=clean(d.get("거래점명")),
                        입금인코드=clean(d.get("입금인코드")), 메모=clean(d.get("메모"))))
    return out


def bank_t_rows(rows):
    if find_table(rows, T_REQ)[0] is None:
        return None
    out = []
    for d in table_dicts(rows, T_REQ):
        if not d.get("거래일시"):
            continue
        try:
            t = to_ts(d["거래일시"])
        except ValueError:
            continue
        out.append(dict(t=t, 구분=clean(d.get("이체구분")), 결과=clean(d.get("처리결과")), 출금계좌=clean(d.get("출금계좌")),
                        은행=clean(d.get("입금은행")), 계좌=clean(d.get("입금계좌")), 수취인=clean(d.get("수취인성명")),
                        금액=num(d.get("처리금액")), 입금금액=num(d.get("입금금액")), 수수료=num(d.get("수수료")),
                        오류금액=num(d.get("오류금액")), 등록일자=clean(d.get("등록일자")), 오류코드=clean(d.get("오류코드")),
                        메모=clean(d.get("거래메모")), 입금인코드=clean(d.get("입금인코드")),
                        출금표시=clean(d.get("출금통장표시내용")), 입금표시=clean(d.get("입금통장표시내용")),
                        부담=clean(d.get("수취인수수료부담여부")), used=False))
    return out


def qkey(q):
    return (q["t"], q["입금"], q["출금"], q["잔액"])


def tkey(t):
    return (t["t"], N(t["수취인"] or "").replace(" ", ""), t["금액"], N(t["계좌"] or "").replace(" ", "").replace("-", ""))


def bank_files(base, prefix):
    return sorted((p for p in base.glob("*.xlsx") if N(p.name).startswith(prefix) and not p.name.startswith("~$")),
                  key=lambda p: N(p.name))


def merge_unique(groups, key, seen=None):
    """파일마다 겹치는 기간을 한 번만 남긴다. 같은 파일 안의 똑같은 행(같은 초에 같은 금액 다계좌 이체)은 그대로 둔다."""
    from collections import Counter
    have = Counter(seen or {})
    out = []
    for rows in groups:
        nth = Counter()
        for r in rows:
            k = key(r); nth[k] += 1
            if nth[k] > have[k]:
                have[k] += 1; out.append(r)
    return out


def load_groups(base, prefix, parse):
    groups = []
    for p in bank_files(base, prefix):
        for _, rows in read_rows(p):
            groups.append(parse(rows) or [])
    return groups


def load_all_q(base):
    return merge_unique(load_groups(base, "조회", bank_q_rows), qkey)


def load_all_t(base):
    return merge_unique(load_groups(base, "이체", bank_t_rows), tkey)


def append_extra(path, cols, rows, to_row):
    """추가분 파일(원본과 같은 열 순서)에 행을 붙인다."""
    if path.exists():
        wb = openpyxl.load_workbook(str(path)); ws = wb.active
    else:
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "sheet 1"; ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
    n = ws.max_row - 1
    for r in sorted(rows, key=lambda x: x["t"]):
        n += 1
        ws.append([n] + to_row(r))
    wb.save(str(path))


def q_to_row(q):
    return [None, q["t"].strftime("%Y.%m.%d %H:%M:%S"), q["적요"], q["입금"], q["출금"], q["내용"], q["잔액"], q["점"],
            q.get("입금인코드"), q.get("메모")]


def t_to_row(t):
    return [None, t["구분"], t["t"].strftime("%Y.%m.%d %H:%M:%S"), t["결과"], t["출금계좌"], t["은행"], t["계좌"], t["수취인"],
            t["금액"], t["입금금액"], t["수수료"], t["오류금액"], t["등록일자"], t["오류코드"], t["메모"], t["입금인코드"],
            t["출금표시"], t["입금표시"], t["부담"]]


# ------------------------------------------------------------------ 홈택스 세금계산서

INV_REQ = ["작성일자", "승인번호", "공급가액", "세액"]


def hometax_rows(rows, title_text):
    """홈택스 목록조회 엑셀 → 계산서 dict 목록. 상호·대표자명 등 겹치는 머리글은 앞=공급자, 뒤=공급받는자."""
    hdr, body = find_table(rows, INV_REQ)
    if hdr is None:
        return None
    idx = {}
    for i, h in enumerate(hdr):
        idx.setdefault(h, []).append(i)

    def col(r, *names, nth=0):
        for name in names:
            pos = idx.get(norm_header(name))
            if pos and len(pos) > nth and pos[nth] < len(r):
                v = r[pos[nth]]
                return clean(v) if v not in ("",) else None
        return None

    def reg(r, side):
        v = col(r, f"{side}사업자등록번호", f"{side}등록번호")
        if v is None:
            v = col(r, "사업자등록번호", "등록번호", nth=0 if side == "공급자" else 1)
        return N(v).strip() if v else None

    out = []
    for r in body:
        appr = col(r, "승인번호")
        if not appr or not re.match(r"^\d{8}-", str(appr)):
            continue
        out.append(dict(
            작성일자=to_date(col(r, "작성일자")), 승인=N(appr).strip(),
            발급일자=col(r, "발급일자"), 전송일자=col(r, "전송일자"),
            공급자번호=reg(r, "공급자"), 공급자상호=col(r, "공급자상호", "상호", nth=0) if "공급자상호" in idx else col(r, "상호", nth=0),
            공급자대표=col(r, "공급자대표자명", "대표자명", nth=0) if "공급자대표자명" in idx else col(r, "대표자명", nth=0),
            공급자주소=col(r, "공급자주소") or col(r, "주소", nth=0),
            받는자번호=reg(r, "공급받는자"),
            받는자상호=col(r, "공급받는자상호") or col(r, "상호", nth=1),
            받는자대표=col(r, "공급받는자대표자명") or col(r, "대표자명", nth=1),
            합계=num(col(r, "합계금액")), 공급가액=num(col(r, "공급가액")), 세액=num(col(r, "세액")),
            분류=col(r, "전자세금계산서분류", "분류"), 종류=col(r, "전자세금계산서종류", "종류"),
            발급유형=col(r, "발급유형"), 비고=col(r, "비고"), 영청=col(r, "영수/청구구분", "영수/청구"),
            공급자메일=col(r, "공급자이메일"), 받는자메일=col(r, "공급받는자이메일1", "공급받는자이메일"),
            품목=col(r, "품목명"), 수량=num_or_none(col(r, "품목수량", "수량")), 단가=num_or_none(col(r, "품목단가", "단가")),
        ))
        d = out[-1]
        if not d["합계"]:
            d["합계"] = d["공급가액"] + d["세액"]
    return out


def side_of(invs, title_text, cfg, buy_regs):
    """매출/매입 판단: 우리 사업자번호(설정) → 표 제목 → 모든 행에 같은 번호가 있는 쪽 → 기존 매입처 번호."""
    me = cfg.get("my_reg")
    sup = {i["공급자번호"] for i in invs}; rec = {i["받는자번호"] for i in invs}
    if me:
        if sup == {me}: return "매출"
        if rec == {me}: return "매입"
    t = N(title_text)
    if "매출" in t and "매입" not in t: return "매출"
    if "매입" in t and "매출" not in t: return "매입"
    if len(invs) >= 2:
        if len(sup) == 1 and len(rec) > 1: return "매출"
        if len(rec) == 1 and len(sup) > 1: return "매입"
    if sup & buy_regs: return "매입"
    if len(sup) == 1 and len(rec) == 1 and sup != rec:
        return None
    return None


def period_label(d):
    return f"{d.year}-{1 if d.month <= 6 else 2}기"


def data_end_row(ws):
    """'합계' 행 바로 위까지가 자료."""
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "합계":
            return r - 1, r
    return ws.max_row, None


def copy_style(src, dst):
    dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format


def append_invoices(path, sheet, new_rows, date_col, key_col):
    """통합 계산서 파일에 행을 더하고 작성일자 순으로 다시 정렬, 집계 시트도 맞춘다."""
    wb = openpyxl.load_workbook(str(path))
    ws = wb[sheet]
    last, total_row = data_end_row(ws)
    width = ws.max_column
    rows = [[ws.cell(r, c).value for c in range(1, width + 1)] for r in range(2, last + 1)]
    style_row = [ws.cell(2, c) for c in range(1, width + 1)]
    styles = [(copy(c.font), copy(c.fill), copy(c.border), copy(c.alignment), c.number_format) for c in style_row]
    total_vals = [ws.cell(total_row, c).value for c in range(1, width + 1)] if total_row else None
    total_styles = [(copy(ws.cell(total_row, c).font), copy(ws.cell(total_row, c).fill), copy(ws.cell(total_row, c).border),
                     copy(ws.cell(total_row, c).alignment), ws.cell(total_row, c).number_format)
                    for c in range(1, width + 1)] if total_row else None
    rows += new_rows

    def sort_key(r):
        try:
            return to_date(r[date_col])
        except Exception:
            return dt.date.max
    rows.sort(key=sort_key)
    if total_row:
        ws.delete_rows(2, ws.max_row)
    else:
        ws.delete_rows(2, ws.max_row)
    for i, r in enumerate(rows, 2):
        r[0] = i - 1
        for c, v in enumerate(r, 1):
            cell = ws.cell(i, c); cell.value = v
            f, fi, b, a, nf = styles[c - 1]
            cell.font = copy(f); cell.fill = copy(fi); cell.border = copy(b); cell.alignment = copy(a); cell.number_format = nf
    end = len(rows) + 1
    if total_vals:
        tr = end + 1
        for c, v in enumerate(total_vals, 1):
            if isinstance(v, str) and v.startswith("="):
                v = re.sub(r"([A-Z]+)2:([A-Z]+)\d+", lambda m: f"{m.group(1)}2:{m.group(2)}{end}", v)
            cell = ws.cell(tr, c); cell.value = v
            f, fi, b, a, nf = total_styles[c - 1]
            cell.font = copy(f); cell.fill = copy(fi); cell.border = copy(b); cell.alignment = copy(a); cell.number_format = nf
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{get_column_letter(width)}{end}"

    # 집계 시트: 없는 연·월·거래처 행을 더한다
    years_months = sorted({(r[2], r[3]) for r in rows if r[2]})
    years = sorted({r[2] for r in rows if r[2]})
    vendors = []
    for r in rows:
        if r[key_col] and r[key_col] not in [v[0] for v in vendors]:
            vendors.append((r[key_col], r))
    for name, keys, nkey in (("월별집계", [list(k) for k in years_months], 2), ("연도별집계", [[y] for y in years], 1)):
        if name in wb.sheetnames:
            ensure_summary_rows(wb[name], keys, nkey, sort=True)
    if "거래처별집계" in wb.sheetnames:
        w = wb["거래처별집계"]
        hdr = [c.value for c in w[1]]
        nkey = sum(1 for c in w[2] if not (isinstance(c.value, str) and c.value.startswith("=")))
        keys = []
        for reg_no, r in vendors:
            if sheet == "매입세금계산서":
                keys.append([reg_no, r[7]][:nkey])
            else:
                keys.append([reg_no, r[7], r[8]][:nkey])
        ensure_summary_rows(w, keys, nkey, sort=False)
    wb.save(str(path))


def ensure_summary_rows(ws, keys, nkey, sort):
    last, total_row = data_end_row(ws)
    width = ws.max_column
    tmpl = [ws.cell(2, c).value for c in range(1, width + 1)]
    sty = [(copy(ws.cell(2, c).font), copy(ws.cell(2, c).fill), copy(ws.cell(2, c).border), copy(ws.cell(2, c).alignment),
            ws.cell(2, c).number_format) for c in range(1, width + 1)]
    existing = [[ws.cell(r, c).value for c in range(1, width + 1)] for r in range(2, last + 1)]
    have = {tuple(r[:1]) if nkey > 1 and not sort else tuple(r[:nkey]) for r in existing}
    added = False
    for k in keys:
        probe = tuple(k[:1]) if nkey > 1 and not sort else tuple(k[:nkey])
        if probe in have:
            continue
        existing.append(list(k) + [None] * (width - len(k))); have.add(probe); added = True
    if not added:
        return
    if sort:
        existing.sort(key=lambda r: tuple(x if x is not None else 0 for x in r[:nkey]))
    total_vals = [ws.cell(total_row, c).value for c in range(1, width + 1)] if total_row else None
    tsty = [(copy(ws.cell(total_row, c).font), copy(ws.cell(total_row, c).fill), copy(ws.cell(total_row, c).border),
             copy(ws.cell(total_row, c).alignment), ws.cell(total_row, c).number_format)
            for c in range(1, width + 1)] if total_row else None
    ws.delete_rows(2, ws.max_row)
    for i, r in enumerate(existing, 2):
        for c in range(1, width + 1):
            v = tmpl[c - 1]
            if isinstance(v, str) and v.startswith("="):
                v = re.sub(r"(?<![A-Z$])([A-Z]{1,2})2(?!\d)", lambda m: f"{m.group(1)}{i}", v)
            else:
                v = r[c - 1]
            cell = ws.cell(i, c); cell.value = v
            f, fi, b, a, nf = sty[c - 1]
            cell.font = copy(f); cell.fill = copy(fi); cell.border = copy(b); cell.alignment = copy(a); cell.number_format = nf
    end = len(existing) + 1
    if total_vals:
        tr = end + 1
        for c, v in enumerate(total_vals, 1):
            if isinstance(v, str) and v.startswith("="):
                v = re.sub(r"([A-Z]+)2:([A-Z]+)\d+", lambda m: f"{m.group(1)}2:{m.group(2)}{end}", v)
            cell = ws.cell(tr, c); cell.value = v
            f, fi, b, a, nf = tsty[c - 1]
            cell.font = copy(f); cell.fill = copy(fi); cell.border = copy(b); cell.alignment = copy(a); cell.number_format = nf


def buy_row(i):
    d = i["작성일자"]
    return [None, d.strftime("%Y-%m-%d"), d.year, d.month, fmt_d(i["발급일자"]), fmt_d(i["전송일자"]), i["공급자번호"],
            i["공급자상호"], i["공급자대표"], i["공급자주소"], i["품목"], i["수량"], i["단가"], i["공급가액"], i["세액"],
            i["합계"], i["승인"], i["분류"], i["종류"], i["발급유형"], i["영청"], i["비고"], i["공급자메일"], period_label(d)]


def sell_row(i):
    d = i["작성일자"]
    return [None, dt.datetime(d.year, d.month, d.day), d.year, d.month, fmt_d(i["발급일자"]), fmt_d(i["전송일자"]),
            i["받는자번호"], i["받는자상호"], i["받는자대표"], i["품목"], i["공급가액"], i["세액"], i["합계"], i["승인"],
            i["종류"], i["발급유형"], i["영청"], i["비고"], i["받는자메일"]]


def fmt_d(v):
    if v is None:
        return None
    try:
        return to_date(v).strftime("%Y-%m-%d")
    except Exception:
        return v


def existing_approvals(path, sheet, col):
    if not path.exists():
        return set(), set()
    ws = openpyxl.load_workbook(str(path), read_only=True)[sheet]
    appr, regs = set(), set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if len(r) > col and r[col]:
            appr.add(N(r[col]).strip())
        if len(r) > 6 and r[6]:
            regs.add(N(r[6]).strip())
    return appr, regs


# ------------------------------------------------------------------ 사용자 손편집 보존


def keyval(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time() else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def last_data_row(ws, start=2):
    r = start
    while r <= ws.max_row and any(ws.cell(r, c).value is not None for c in range(1, ws.max_column + 1)):
        r += 1
    return r - 1


def copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(c.row, c.column, c.value)
            if c.has_style:
                copy_style(c, d)
    for k, cd in src.column_dimensions.items():
        dst.column_dimensions[k].width = cd.width
    dst.freeze_panes = src.freeze_panes
    if src.auto_filter.ref:
        dst.auto_filter.ref = src.auto_filter.ref


def moved(v, old_row, new_row, col):
    """손으로 넣은 수식은 행이 옮겨진 만큼 참조도 옮긴다(=E8+E6 → 새 위치 기준)."""
    if not (isinstance(v, str) and v.startswith("=")) or old_row == new_row:
        return v
    from openpyxl.formula.translate import Translator
    letter = get_column_letter(col)
    try:
        return Translator(v, origin=f"{letter}{old_row}").translate_formula(f"{letter}{new_row}")
    except Exception:
        return v


def merge_into_existing(new_wb, out_path, generated, always_new):
    """기존 통합 파일을 열어 자동 생성 시트만 새 내용으로 바꾼다.
    손으로 더한 열(자동 열 오른쪽)과 표 아래 메모, 직접 만든 시트, 지운 시트 상태는 그대로 둔다."""
    if not out_path.exists():
        return new_wb
    old = openpyxl.load_workbook(str(out_path))
    for name in generated:
        if name not in new_wb.sheetnames:
            continue
        if name not in old.sheetnames and name not in always_new:
            continue  # 사용자가 지운 시트는 다시 만들지 않는다
        src = new_wb[name]
        width = len([c for c in src[1] if c.value is not None])
        start = 3 if name.startswith("가수금") else 2
        kcols = list(range(2, min(width, 9) + 1))   # No 를 뺀 앞쪽 열(거래일자·시간·금액…)로 같은 행을 찾는다
        extras, trailing, extra_hdr, widths = {}, [], [], {}
        if name in old.sheetnames:
            ow = old[name]
            oend = last_data_row(ow, start)
            extra_cols = list(range(width + 1, ow.max_column + 1))
            for c in extra_cols:
                if ow.cell(1, c).value not in (None, ""):
                    extra_hdr.append(ow.cell(1, c))
                letter = get_column_letter(c)
                if letter in ow.column_dimensions:
                    widths[letter] = ow.column_dimensions[letter].width
            for r in range(start, oend + 1):
                vals = {c: ow.cell(r, c) for c in extra_cols if ow.cell(r, c).value is not None}
                if vals:
                    key = tuple(keyval(ow.cell(r, c).value) for c in kcols)
                    extras.setdefault(key, []).append((r, vals))
            if last_data_row(src, start) >= src.max_row:   # 자동 내용이 표 아래에 없을 때만 손메모를 옮긴다
                for r in range(oend + 1, ow.max_row + 1):
                    cells = [ow.cell(r, c) for c in range(1, ow.max_column + 1)]
                    if any(c.value is not None for c in cells):
                        trailing.append((r - oend, cells))
            idx = old.sheetnames.index(name)
            del old[name]
            dst = old.create_sheet(name, idx)
        else:
            dst = old.create_sheet(name, 0)
        copy_sheet(src, dst)
        for h in extra_hdr:
            copy_style(h, dst.cell(1, h.column, h.value))
        for letter, w in widths.items():
            dst.column_dimensions[letter].width = w
        nend = last_data_row(dst, start)
        for r in range(start, nend + 1):
            lst = extras.get(tuple(keyval(dst.cell(r, c).value) for c in kcols))
            if lst:
                orow, vals = lst.pop(0)
                for c, cell in vals.items():
                    copy_style(cell, dst.cell(r, c, moved(cell.value, orow, r, c)))
        for off, cells in trailing:
            for c in cells:
                if c.value is not None:
                    copy_style(c, dst.cell(nend + off, c.column, moved(c.value, c.row, nend + off, c.column)))
    return old


# ------------------------------------------------------------------ 이번 정리 시트

PINK = PatternFill("solid", fgColor="FFC7CE"); PINK_F = Font(color="9C0006")
GREEN = PatternFill("solid", fgColor="C6EFCE"); GREEN_F = Font(color="006100")
GRAY = PatternFill("solid", fgColor="F2F2F2"); GRAY_F = Font(color="7F7F7F")
HDR_F = Font(bold=True, color="FFFFFF"); HDR_FILL = PatternFill("solid", fgColor="1F4E78")


DAILY_FILE = "일용직 내역.xlsx"
COMPANY_WORDS = re.compile(r"주식회사|\(주\)|㈜|\(유\)|상사|건재|건설|산업|공사|센터|조명|화학|자재|철물|목재|설비|전기|차|국세|세|보험|공단|은행|카드")


def daily_names(base):
    """일용직 내역 파일에 있는 이름(지급명세서 성명, 통합내역 수취인)."""
    names = set()
    p = base / DAILY_FILE
    if not p.exists():
        return names
    try:
        wb = openpyxl.load_workbook(str(p), data_only=True, read_only=False)
    except Exception:
        return names
    for ws in wb.worksheets:
        if "일용" not in N(ws.title) and N(ws.title) != "통합내역":   # 같은 파일의 다른 시트(이체결과 사본 등)는 쓰지 않는다
            continue
        for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True)):
            hdr = [norm_header(v) for v in row]
            col = next((hdr.index(h) for h in ("성명", "수취인성명") if h in hdr), None)
            if col is None:
                continue
            for r in ws.iter_rows(min_row=i + 2, values_only=True):
                v = r[col] if col < len(r) else None
                if isinstance(v, str) and 2 <= len(v.strip()) <= 10 and not COMPANY_WORDS.search(v):
                    names.add(v.strip())
            break
    return names


def daily_reason(q, names):
    """일용직 판정(메모리 규칙): 명단에 있거나, 메모에 일용/일당, 또는 개인에게 15~20만원."""
    if q["출금"] <= 0:
        return None
    t = q.get("T")
    who = (t and t["수취인"]) or q["내용"] or ""
    who = who.strip()
    memo = " ".join(str(x) for x in (q["내용"], t and t["메모"], t and t["입금표시"]) if x)
    if who in names:
        return "일용직 명단"
    if re.search(r"일용|일당", memo):
        return "메모(일용/일당)"
    amt = t["금액"] if t else q["출금"]
    if 150000 <= amt <= 200000 and re.fullmatch(r"[가-힣]{2,4}|[A-Z]{3,12}", who) and not COMPANY_WORDS.search(who):
        return "개인 15~20만원"
    return None


def category(q, names=frozenset()):
    if "refund" in q: return "오송금 반환", None
    if q.get("tax"): return "세금·공과금", q["tax"]
    if q.get("pay"): return "급여(김한승)", None
    if q.get("loan"): return "가수금(김한승)", None
    if q["적요"] == "신한체": return "체크카드", None
    if q["입금"] > 0: return "입금내역", None
    reason = daily_reason(q, names)
    if reason and not q.get("inv"):
        return "일용직", reason
    return "통합내역(거래처 출금)", None


def needs_invoice(sheet, q):
    """세금계산서가 있어야 정상인 거래인가 — 미매칭이면 분홍으로 표시."""
    if sheet == "통합내역(거래처 출금)":
        return not str(q["내용"] or "").upper().startswith("SMS")
    if sheet == "입금내역":
        return True
    return False


def report_sheet(wb, g, new_q, new_inv, summary, names):
    ws = wb.create_sheet(REPORT_SHEET, 0)
    ws.append([f"이번 정리 — {dt.datetime.now():%Y-%m-%d %H:%M}"])
    ws["A1"].font = Font(bold=True, size=14)
    for line in summary:
        ws.append([line])
    ws.append([])
    ws.append(["새 은행 거래 (분홍 = 세금계산서와 맞춰지지 않음)"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    hdr = ["거래일시", "입금액", "출금액", "내용(통장표시)", "수취인", "거래메모", "분류", "세부", "계산서 매칭", "계산서 상호", "계산서 합계", "계산서 작성일자"]
    ws.append(hdr); hr = ws.max_row
    for c in ws[hr]:
        c.font = HDR_F; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center")
    newkeys = {qkey(q) for q in new_q}
    rows = [q for q in g["Q"] if qkey(q) in newkeys]
    rows.sort(key=lambda q: q["t"], reverse=True)
    for q in rows:
        sheet, sub = category(q, names)
        t = q.get("T")
        if q["입금"] > 0:
            inv, invs, mt = q.get("sinv"), q.get("sinvs"), q.get("sinvtype")
            nm = (inv["상호"] or inv["대표"]) if inv else None
        else:
            inv, invs, mt = q.get("inv"), q.get("invs"), q.get("invtype")
            nm = inv["상호"] if inv else None
        if invs:
            nm = f"{nm} 외 {len(invs) - 1}건"; amt = sum(i["합계"] for i in invs); d = ", ".join(i["d"].strftime("%m/%d") for i in invs)
        else:
            amt = inv["합계"] if inv else None; d = inv["d"] if inv else None
        if not mt and str(q["내용"] or "").upper().startswith("SMS"):
            mt = "은행비용(SMS)"
        unsettled = not mt and sheet.startswith("통합내역") and any(
            v in f'{q["내용"]}{t and t["수취인"] or ""}' for v in g.get("PINK_VENDORS", []))
        ws.append([q["t"], q["입금"] or None, q["출금"] or None, q["내용"], t and t["수취인"], t and t["메모"], sheet, sub,
                   "미정산" if unsettled else mt, nm, amt, d])
        r = ws.max_row
        ws.cell(r, 1).number_format = "yyyy-mm-dd hh:mm"; ws.cell(r, 12).number_format = "yyyy-mm-dd"
        for c in (2, 3, 11):
            ws.cell(r, c).number_format = "#,##0"
        if sheet == "오송금 반환":
            fill, font = GRAY, GRAY_F
        elif mt:
            fill, font = GREEN, GREEN_F
        elif needs_invoice(sheet, q):
            fill, font = PINK, PINK_F
        else:
            fill, font = GRAY, GRAY_F
        for c in ws[r]:
            c.fill = fill; c.font = font
    if not rows:
        ws.append(["(새 은행 거래 없음)"])
    ws.append([])
    ws.append(["새 세금계산서 (분홍 = 은행 거래와 맞춰지지 않음)"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["구분", "작성일자", "상호", "대표자명", "품목명", "합계금액", "승인번호", "은행 매칭", "은행 거래일자", "은행 금액", "은행 내용"])
    for c in ws[ws.max_row]:
        c.font = HDR_F; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center")
    appr_new = {i["승인"] for i in new_inv}
    inv_rows = []
    for kind, lst in (("매입", g["INV"]), ("매출", g["SINV"])):
        for i in lst:
            if N(i["승인"] or "") in appr_new:
                inv_rows.append((kind, i))
    for kind, i in sorted(inv_rows, key=lambda x: x[1]["d"], reverse=True):
        b = i["bank"] if i["bank"] not in (None, "X") else None
        if i.get("cancel"): mt = "취소/정정 상계"
        elif i.get("vendor_ok"): mt = "거래처 합계 일치"
        elif "split" in i: mt = "분할 일치"
        elif "merged" in i: mt = "합산 일치"
        elif b: mt = b.get("sinvtype" if kind == "매출" else "invtype")
        elif i["합계"] < 0: mt = "수정/취소(음수)"
        else: mt = None
        bd = b["t"] if b else None
        ba = (b["입금"] if kind == "매출" else b["출금"]) if b else None
        ws.append([kind, i["d"], i["상호"], i["대표"], i["품목"], i["합계"], i["승인"], mt, bd, ba, b["내용"] if b else None])
        r = ws.max_row
        ws.cell(r, 2).number_format = "yyyy-mm-dd"; ws.cell(r, 9).number_format = "yyyy-mm-dd"
        for c in (6, 10):
            ws.cell(r, c).number_format = "#,##0"
        fill, font = (GREEN, GREEN_F) if mt else (PINK, PINK_F)
        for c in ws[r]:
            c.fill = fill; c.font = font
    if not inv_rows:
        ws.append(["(새 세금계산서 없음)"])
    for i, w in enumerate([19, 13, 13, 22, 14, 18, 18, 14, 16, 24, 13, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.tabSelected = True
    return ws


# ------------------------------------------------------------------ 현장별 합계

SITE_SHEET = "현장별 합계"
SITE_HDR = "현장"
SITE_ROWS = 100          # 현장 이름을 몇 개까지 모을지
SITE_FILL = PatternFill("solid", fgColor="FFF2CC")
# (시트, 표시 이름, 더할 금액 열, 뺄 금액 열(카드 취소·환불), 날짜 열)
SITE_SOURCES = [
    ("통합내역", "이체(거래처·일용직)", "출금액", None, "거래일자"),
    ("체크카드", "체크카드", "사용금액", "입금(취소/환불)", "거래일자"),
    ("세금·공과금", "세금·공과금", "출금액", None, "거래일자"),
    ("급여(김한승)", "급여", "출금액", None, "거래일자"),
]


def ensure_site_column(ws):
    """시트 오른쪽 끝(손으로 더한 열 다음)에 「현장」 열을 둔다. 써 둔 값은 merge_into_existing 이 옮겨 준다."""
    hdr = [c.value for c in ws[1]]
    if SITE_HDR in hdr:
        return hdr.index(SITE_HDR) + 1
    col = len(hdr) + 1
    while col > 1 and ws.cell(1, col - 1).value in (None, ""):
        col -= 1
    c = ws.cell(1, col, SITE_HDR)
    c.font = Font(bold=True, color="7F6000"); c.fill = PatternFill("solid", fgColor="FFD966")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(col)].width = 18
    return col


def build_site_sheet(wb):
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.datavalidation import DataValidation
    first, end = 4, 4 + SITE_ROWS - 1
    srcs = []
    for name, label, plus, minus, date in SITE_SOURCES:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        col = ensure_site_column(ws)
        hdr = [c.value for c in ws[1]]
        if plus not in hdr:
            continue
        start = 3 if name.startswith("가수금") else 2
        last = max(ws.max_row, 2) + 500        # 앞으로 늘어날 행까지
        q = f"'{name}'!"
        letter = lambda h: get_column_letter(hdr.index(h) + 1)
        rng = f"{q}${get_column_letter(col)}${start}:${get_column_letter(col)}${last}"
        part = lambda h: f"{q}${letter(h)}${start}:${letter(h)}${last}"
        srcs.append(dict(ws=ws, col=col, label=label, rng=rng, plus=part(plus),
                         minus=part(minus) if minus and minus in hdr else None,
                         date=part(date) if date in hdr else None, start=start, last=last))
    if not srcs:
        return
    if SITE_SHEET in wb.sheetnames:
        idx = wb.sheetnames.index(SITE_SHEET); del wb[SITE_SHEET]
    else:
        idx = wb.sheetnames.index(srcs[0]["ws"].title) + 1
    ws = wb.create_sheet(SITE_SHEET, idx)
    ws["A1"] = "현장별 출금 합계 — " + ", ".join(f"「{s['ws'].title}」" for s in srcs) + " 시트의 「현장」 열에 이름을 쓰면 여기에 모입니다"
    ws["A1"].font = Font(bold=True, size=13)
    heads = ["현장", "건수"] + [s["label"] for s in srcs] + ["총합", "마지막 거래일"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(3, i, h); c.font = HDR_F; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", wrap_text=True)
    ncol = len(srcs)
    tot_col = 3 + ncol; date_col = tot_col + 1
    ws["A2"] = "전체"; ws["A2"].font = Font(bold=True)
    for c in range(2, tot_col + 1):
        L = get_column_letter(c)
        ws.cell(2, c, f"=SUM({L}{first}:{L}{end})").font = Font(bold=True)
        ws.cell(2, c).number_format = "#,##0"
    for r in range(first, end + 1):
        # 각 시트 현장 열에서 아직 위에 없는 이름을 차례로 꺼낸다(배열 수식, 앞 시트부터)
        f = '""'
        for s in reversed(srcs):
            rg = s["rng"]
            f = f'IFERROR(INDEX({rg},MATCH(0,COUNTIF($A${first - 1}:A{r - 1},{rg})+({rg}=""),0)),{f})'
        ws[f"A{r}"] = ArrayFormula(f"A{r}", "=" + f)
        blank = f'A{r}=""'
        ws.cell(r, 2, f'=IF({blank},"",' + "+".join(f"COUNTIF({s['rng']},A{r})" for s in srcs) + ")")
        for i, s in enumerate(srcs):
            expr = f"SUMIFS({s['plus']},{s['rng']},A{r})"
            if s["minus"]:
                expr += f"-SUMIFS({s['minus']},{s['rng']},A{r})"
            ws.cell(r, 3 + i, f'=IF({blank},"",{expr})')
        ws.cell(r, tot_col, f'=IF({blank},"",SUM({get_column_letter(3)}{r}:{get_column_letter(2 + ncol)}{r}))')
        dates = ",".join(f"_xlfn.MAXIFS({s['date']},{s['rng']},A{r})" for s in srcs if s["date"])
        ws.cell(r, date_col, f'=IF({blank},"",MAX({dates}))')
        for c in range(2, tot_col + 1):
            ws.cell(r, c).number_format = "#,##0"
        ws.cell(r, tot_col).font = Font(bold=True)
        ws.cell(r, date_col).number_format = "yyyy-mm-dd"
    widths = [22, 7] + [15] * ncol + [16, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "B4"
    # 현장 칸마다 드롭다운(이미 쓴 현장 이름). 새 이름도 쓸 수 있게 오류는 띄우지 않는다
    for s in srcs:
        src = s["ws"]; L = get_column_letter(s["col"])
        src.data_validations.dataValidation = [dv for dv in src.data_validations.dataValidation
                                               if not str(dv.sqref).startswith(L)]
        dv = DataValidation(type="list", formula1=f"='{SITE_SHEET}'!$A${first}:$A${end}", allow_blank=True,
                            showErrorMessage=False)
        dv.add(f"{L}{s['start']}:{L}{s['last']}")
        src.add_data_validation(dv)
        end_row = last_data_row(src, s["start"])
        for r in range(s["start"], end_row + 1):
            src.cell(r, s["col"]).fill = SITE_FILL


# ------------------------------------------------------------------ 실행


def backup(base, names):
    dst = base / "백업" / dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    dst.mkdir(parents=True, exist_ok=True)
    for n in names:
        p = base / n
        if p.exists():
            shutil.copy2(p, dst / n)
    # 오래된 백업은 최근 20개만 남긴다
    olds = sorted((p for p in (base / "백업").iterdir() if p.is_dir()), key=lambda p: p.name)
    for p in olds[:-20]:
        shutil.rmtree(p, ignore_errors=True)
    return dst


def run(paths):
    base = base_dir()
    if not base.is_dir():
        raise SystemExit(f"은행 폴더가 없습니다: {base}")
    cfg = load_config()
    out_path = base / OUT_NAME
    bk = backup(base, [OUT_NAME, BUY_NAME, SELL_NAME, Q_EXTRA, T_EXTRA])
    lines = []

    old_q = load_all_q(base); old_t = load_all_t(base)
    from collections import Counter
    qseen = Counter(qkey(q) for q in old_q); tseen = Counter(tkey(t) for t in old_t)
    new_q, new_t, buy_new, sell_new, skipped = [], [], [], [], []
    buy_appr, buy_regs = existing_approvals(base / BUY_NAME, "매입세금계산서", 16)
    sell_appr, _ = existing_approvals(base / SELL_NAME, "매출세금계산서", 13)

    for p in paths:
        p = Path(p)
        name = N(p.name)
        try:
            sheets = read_rows(p)
        except Exception as e:
            skipped.append(f"{name}: 열지 못함 ({e})"); continue
        kind = None
        for _, rows in sheets:
            title = " ".join(N(v) for r in rows[:6] for v in r if v) + " " + name
            qs = bank_q_rows(rows)
            if qs:
                kind = "조회"
                got = merge_unique([qs], qkey, qseen)
                new_q += got; qseen.update(qkey(q) for q in got)
                continue
            ts = bank_t_rows(rows)
            if ts:
                kind = "이체결과"
                got = merge_unique([ts], tkey, tseen)
                new_t += got; tseen.update(tkey(t) for t in got)
                continue
            invs = hometax_rows(rows, title)
            if invs:
                side = side_of(invs, title, cfg, buy_regs)
                if side is None:
                    skipped.append(f"{name}: 매출인지 매입인지 알 수 없음"); kind = "?"; continue
                kind = f"세금계산서({side})"
                me = (invs[0]["공급자번호"] if side == "매출" else invs[0]["받는자번호"])
                if me and len(invs) >= 2 and not cfg.get("my_reg"):
                    cfg["my_reg"] = me; save_config(cfg)
                for i in invs:
                    if side == "매입" and i["승인"] not in buy_appr:
                        buy_appr.add(i["승인"]); buy_new.append(i)
                    elif side == "매출" and i["승인"] not in sell_appr:
                        sell_appr.add(i["승인"]); sell_new.append(i)
        if kind is None:
            skipped.append(f"{name}: 신한 조회·이체결과나 홈택스 세금계산서 양식이 아님")

    if new_q:
        append_extra(base / Q_EXTRA, Q_COLS, new_q, q_to_row)
    if new_t:
        append_extra(base / T_EXTRA, T_COLS, new_t, t_to_row)
    if buy_new:
        append_invoices(base / BUY_NAME, "매입세금계산서", [buy_row(i) for i in buy_new], 1, 6)
    if sell_new:
        append_invoices(base / SELL_NAME, "매출세금계산서", [sell_row(i) for i in sell_new], 1, 6)

    lines.append(f"추가: 은행 조회 {len(new_q)}건 · 이체결과 {len(new_t)}건 · 매입계산서 {len(buy_new)}건 · 매출계산서 {len(sell_new)}건")
    for s in skipped:
        lines.append("건너뜀: " + s)

    # 통합 다시 만들기 (분류·매칭 규칙은 integrate.py)
    import contextlib, io
    with contextlib.redirect_stdout(io.StringIO()):   # integrate.py 의 진행 기록은 결과 창에 보이지 않게
        g = runpy.run_path(str(HERE / "integrate.py"), init_globals=dict(
            BASE=str(base), LOAD_Q=lambda: load_all_q(base), LOAD_T=lambda: load_all_t(base)))
    new_wb = g["wb"]
    generated = list(new_wb.sheetnames)
    unmatched_new = 0
    names = daily_names(base)
    newkeys = {qkey(q) for q in new_q}
    for q in g["Q"]:
        if qkey(q) in newkeys:
            sheet, _ = category(q, names)
            mt = q.get("sinvtype") if q["입금"] > 0 else q.get("invtype")
            if not mt and needs_invoice(sheet, q):
                unmatched_new += 1
    lines.append(f"새 은행 거래 중 계산서 미매칭(분홍) {unmatched_new}건")
    if paths:
        report_sheet(new_wb, g, new_q, buy_new + sell_new, lines, names)
        generated.append(REPORT_SHEET)
    # 파일 없이 다시 정리할 때는 지난번 「이번 정리」 시트를 그대로 둔다
    final = merge_into_existing(new_wb, out_path, generated, always_new={REPORT_SHEET})
    build_site_sheet(final)
    if REPORT_SHEET in final.sheetnames:
        final.active = final.sheetnames.index(REPORT_SHEET)
        for w in final.worksheets:
            w.sheet_view.tabSelected = (w.title == REPORT_SHEET)
    final.save(str(out_path))
    lines.append(f"백업: {bk.relative_to(base)}")
    return out_path, lines


def target_names():
    """버튼이 열려 있는 통합 파일을 찾아 닫을 수 있게 이름·경로를 NFC/NFD 둘 다 알려준다."""
    p = base_dir() / OUT_NAME
    out = []
    for form in ("NFC", "NFD"):
        out += [unicodedata.normalize(form, p.name), unicodedata.normalize(form, str(p))]
    print("\n".join(dict.fromkeys(out)))


def main(argv):
    if argv[:1] == ["--target-names"]:
        return target_names()
    paths = [a for a in argv if a.strip()]
    if len(paths) == 1 and "\n" in paths[0]:
        paths = [p for p in paths[0].split("\n") if p.strip()]
    paths = [hfs_to_posix(p.strip()) for p in paths]
    try:
        out, lines = run(paths)
        print(f"OPEN:{out}")
        print("\n".join(lines))
    except SystemExit as e:
        print("ERROR:" + str(e))
    except Exception as e:
        log = HOME / "error.log"
        HOME.mkdir(parents=True, exist_ok=True)
        log.write_text(traceback.format_exc(), "utf-8")
        print(f"ERROR:정리 중 오류가 났습니다: {e}\n자세한 내용: {log}")


def hfs_to_posix(p):
    if p.startswith("/") or ":" not in p:
        return p
    parts = p.split(":")
    if parts[0] in ("Macintosh HD",):
        return "/" + "/".join(parts[1:])
    return "/Volumes/" + "/".join(parts)


if __name__ == "__main__":
    main(sys.argv[1:])
