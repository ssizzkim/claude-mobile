# 은행내역 통합 (조회 + 이체결과 + 매입/매출세금계산서 + 급여명세) — 분류·매칭 규칙
# 원래 「은행내역 통합_스크립트.py」(2026-09-18). 파일 읽기와 저장만 banksort.py 로 옮겼다.
import openpyxl, glob, os, unicodedata, warnings, datetime as dt, re
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from itertools import combinations
from collections import OrderedDict
from copy import copy
warnings.filterwarnings("ignore")
N = lambda s: unicodedata.normalize("NFC", s)
# banksort.py 가 BASE(은행 폴더), LOAD_Q, LOAD_T 를 넣어 실행한다.
files = sorted((f for f in glob.glob(os.path.join(BASE, "*.xlsx")) if not os.path.basename(f).startswith("~$")), key=N)
Q = LOAD_Q()
T = LOAD_T()
Q.sort(key=lambda x:x["t"]); T.sort(key=lambda x:x["t"])

# ---------- 조회 ↔ 이체결과 매칭 ----------
matched=0
def name_ok(q,t):
    a=(q["내용"] or "").replace(" ",""); b=(t["수취인"] or "").replace(" ","")
    return bool(a) and bool(b) and (a.startswith(b[:3]) or b.startswith(a[:3]))
for mode in (0,1,2):
    for q in Q:
        if q["출금"]<=0 or "T" in q: continue
        best=None
        for t in T:
            if t["used"]: continue
            if mode<2 and q["출금"]!=t["금액"]+t["수수료"]: continue
            if mode==2 and q["출금"]!=t["금액"]: continue
            if mode==0 and not name_ok(q,t): continue
            d=abs((q["t"]-t["t"]).total_seconds())
            if d<=300 and (best is None or d<best[0]): best=(d,t)
        if best:
            best[1]["used"]=True; q["T"]=best[1]; matched+=1
unmatched_T=[t for t in T if not t["used"]]
print("조회",len(Q),"이체",len(T),"매칭",matched,"미매칭이체",len(unmatched_T))

# ---------- 분류 ----------
TAXRULES=[("국세",r"국세"),("지방세/세외수입",r"서울시세입금|서울특별시 세|서울세외|E-tax|스마트로-세외|세외수입|토르크서울주민"),
          ("4대보험",r"사회보험합산|건강보험|국민건강보험공단"),("세무대행",r"이정섭"),("세금납부(주식회사토르크)",r"^FB이체 주식회사토르크")]
def taxcat(q):
    t=q.get("T"); txt=" ".join(str(x) for x in (q["적요"],q["내용"],t and t["수취인"],t and t["입금표시"]) if x)
    for c,p in TAXRULES:
        if re.search(p,txt): return c
    return None
for q in Q: q["tax"]=taxcat(q)
def is_pay(q):
    if q["출금"]<=0 or q["tax"]: return False
    t=q.get("T"); txt=" ".join(str(x) for x in (q["내용"],t and t["수취인"],t and t["메모"],t and t["입금표시"]) if x)
    if re.search(r"대여|반환|상환|가수금|한전",txt): return False
    if q["적요"]=="BZ급여" or re.search(r"급여|월급|\d급",txt): return True
    return bool(t) and t["수취인"]=="김한승" and 2600000<=q["출금"]<=2750000
for q in Q: q["pay"]=is_pay(q)
def is_loan(q):
    if q["tax"] or q["pay"]: return False
    t=q.get("T"); txt=" ".join(str(x) for x in (q["내용"],t and t["수취인"],t and t["메모"],t and t["입금표시"]) if x)
    if "한전" in txt: return False
    if q["입금"]>0: return q["내용"].startswith("김한승")
    return bool(t) and t["수취인"]=="김한승" or q["내용"].startswith("김한승")
for q in Q: q["loan"]=is_loan(q)
QT=[q for q in Q if q["tax"]]; QP=[q for q in Q if q["pay"]]; QL=[q for q in Q if q["loan"]]
QI=[q for q in Q if q["입금"]>0 and q["적요"]!="신한체" and not q["tax"] and not q["pay"] and not q["loan"]]
QM=[q for q in Q if q["입금"]<=0 and q["적요"]!="신한체" and not q["tax"] and not q["pay"] and not q["loan"]]
QC=[q for q in Q if q["적요"]=="신한체" and not q["tax"]]

# ---------- 매입세금계산서 ----------
INV=[]
for f in files:
    if "매입세금계산서" not in N(f): continue
    wsI=openpyxl.load_workbook(f)["매입세금계산서"]
    for r in wsI.iter_rows(min_row=2,values_only=True):
        if isinstance(r[0],(int,float)) and r[1]:
            INV.append(dict(d=dt.datetime.strptime(str(r[1])[:10],"%Y-%m-%d").date(),발급=r[4],사업자=r[6],상호=r[7],대표=r[8],품목=r[10],공급가액=r[13],세액=r[14],합계=r[15],승인=r[16],종류=r[17],구분=r[18],영청=r[20],비고=r[21],bank=None))
def amt_of(q):
    t=q.get("T"); return t["금액"] if t else q["출금"]
def nm(x):
    x=str(x or ""); x=re.sub(r"\(주\)|주식회사|㈜|\(유\)|유한회사|\(합\)|\s|\.|,","",x)
    x=re.sub(r"\(.*?\)|\(.*$","",x); return x
# 수동 별칭: (은행 내용/수취인에 포함된 이름, 계산서 상호에 포함된 이름)
ALIASES=[("김옥순","옥순클린")]
def name_match(q,inv):
    t=q.get("T"); cands=[nm(q["내용"])]+([nm(t["수취인"]),nm(t["입금표시"])] if t else [])
    sup=nm(inv["상호"]); rep=nm(inv["대표"])
    for a,b in ALIASES:
        if any(a in c for c in cands) and b in sup: return True
    for c in cands:
        if len(c)>=2 and len(sup)>=2 and (sup.startswith(c) or c.startswith(sup) or (len(c)>=3 and c in sup)): return True
        if rep and len(rep)>=2 and c.startswith(rep) and t and rep in nm(t["수취인"]): return True
    return False
QX=QM+QT   # 계산서 매칭 대상: 통합내역 + 세금·공과금
def pass_date(need_name,label="날짜+금액 일치"):
    for q in QX:
        if q["출금"]<=0 or "inv" in q: continue
        for inv in INV:
            if inv["bank"] is None and inv["d"]==q["t"].date() and inv["합계"]==amt_of(q) and (not need_name or name_match(q,inv)):
                inv["bank"]=q; q["inv"]=inv; q["invtype"]=label; break
def pass_name():
    pairs=[]
    for inv in INV:
        if inv["bank"]: continue
        for q in QX:
            if q["출금"]<=0 or "inv" in q: continue
            if inv["합계"]!=amt_of(q) or not name_match(q,inv): continue
            pairs.append((abs((q["t"].date()-inv["d"]).days),id(inv),id(q),inv,q))
    pairs.sort(key=lambda p:p[0])
    for d,_,_,inv,q in pairs:
        if inv["bank"] or "inv" in q: continue
        inv["bank"]=q; q["inv"]=inv; q["invtype"]="금액+상호 일치"
# 1차: 날짜+금액+상호 → 2차: 금액+상호(최근접) → 1b: 날짜+금액만
pass_date(True); pass_name(); pass_date(False)
# 3차: 분할납부 — 같은 거래처 미매칭 출금 2~10건 합 = 계산서 금액
def subset_sum(target,items,key):
    dp={0:()}
    for it in items:
        a=key(it); new={}
        for ssum,combo in dp.items():
            ns=ssum+a
            if ns<=target and ns not in dp and ns not in new and len(combo)<10: new[ns]=combo+(it,)
        dp.update(new)
        if target in dp: return dp[target]
    return None
split_n=0
for inv in INV:
    if inv["bank"]: continue
    cands=[q for q in QX if q["출금"]>0 and "inv" not in q and name_match(q,inv) and -60<=(q["t"].date()-inv["d"]).days<=180 and amt_of(q)<inv["합계"]]
    cands.sort(key=lambda q:abs((q["t"].date()-inv["d"]).days)); cands=cands[:20]
    found=subset_sum(inv["합계"],cands,amt_of)
    if found and len(found)>=2:
        inv["bank"]=found[0]; inv["split"]=found; split_n+=1
        for q in found: q["inv"]=inv; q["invtype"]="분할 일치"
print("분할 매칭 계산서",split_n)
# 4차: 합산납부 — 출금 1건 = 같은 거래처 미매칭 계산서 2~10장 합
merge_n=0
for q in QX:
    if q["출금"]<=0 or "inv" in q: continue
    amt=amt_of(q)
    cands=[inv for inv in INV if not inv["bank"] and name_match(q,inv) and -180<=(inv["d"]-q["t"].date()).days<=60 and 0<inv["합계"]<amt]
    cands.sort(key=lambda inv:abs((inv["d"]-q["t"].date()).days)); cands=cands[:20]
    found=subset_sum(amt,cands,lambda i:i["합계"])
    if found and len(found)>=2:
        q["inv"]=found[0]; q["invtype"]="합산 일치"; q["invs"]=found; merge_n+=1
        for inv in found: inv["bank"]=q; inv["merged"]=found
print("합산 매칭 출금",merge_n)
def vkey_p(inv): return nm(inv["대표"]) or nm(inv["상호"])
VP={}
for inv in INV:
    if inv["합계"]>0: VP.setdefault(vkey_p(inv),[]).append(inv)
vendor_ok_p=0
for k,invs in VP.items():
    if all(i["bank"] for i in invs): continue
    outs=[q for q in QX if q["출금"]>0 and any(name_match(q,i) for i in invs) and ("inv" not in q or q["inv"] in invs or (q.get("invs") and all(x in invs for x in q["invs"])))]
    if outs and sum(i["합계"] for i in invs)==sum(amt_of(q) for q in outs):
        for i in invs:
            if not i["bank"]: i["bank"]=outs[0]; i["vendor_ok"]=True
        for q in outs:
            if "inv" not in q: q["inv"]=invs[0]; q["invtype"]="거래처 합계 일치"
        vendor_ok_p+=1
print("매입 거래처합계 일치 거래처",vendor_ok_p)
print("세금계산서",len(INV),"매칭",sum(1 for i in INV if i["bank"]))

# ---------- 매출세금계산서 ↔ 입금 매칭 ----------
SINV=[]
for f in files:
    if "매출세금계산서" not in N(f): continue
    wsS=openpyxl.load_workbook(f)["매출세금계산서"]
    for r in wsS.iter_rows(min_row=2,values_only=True):
        if isinstance(r[0],(int,float)) and r[1]:
            d=r[1].date() if hasattr(r[1],"date") else dt.datetime.strptime(str(r[1])[:10],"%Y-%m-%d").date()
            SINV.append(dict(d=d,발급=r[4],등록=r[6],상호=(r[7] or "").strip(),대표=(r[8] or "").strip(),품목=r[9],공급가액=r[10],세액=r[11],합계=r[12],승인=r[13],영청=r[16],bank=None))
SALIASES=[("DGIST","대구경북과학기술원"),("조종분","테이블아트"),("박효민","채정원")]   # (입금자, 계산서 상호)
def sname_match(q,inv):
    c=nm(q["내용"]); sup=nm(inv["상호"]); rep=nm(inv["대표"])
    for a,b in SALIASES:
        if a.upper() in c.upper() and (b in sup or b in rep): return True
    if len(c)>=2 and len(sup)>=2 and (sup.startswith(c) or c.startswith(sup) or (len(c)>=3 and c in sup)): return True
    if len(c)>=2 and len(rep)>=2 and (rep.startswith(c) or c.startswith(rep)): return True
    return False
# 취소/정정 상계: 음수 계산서 ↔ 같은 등록번호·같은 금액의 이전 양수 계산서
for neg in sorted([i for i in SINV if i["합계"]<0],key=lambda i:i["d"]):
    c=[i for i in SINV if i["합계"]==-neg["합계"] and i["등록"]==neg["등록"] and i["d"]<=neg["d"] and not i.get("cancel")]
    if c:
        pos=max(c,key=lambda i:i["d"]); pos["cancel"]=neg; neg["cancel"]=pos
for i in SINV:
    if i.get("cancel"): i["bank"]="X"
QIx=[q for q in QI]
for need_name in (True,):
    for q in QIx:
        if "sinv" in q: continue
        for inv in SINV:
            if inv["bank"] is None and inv["합계"]>0 and inv["d"]==q["t"].date() and inv["합계"]==q["입금"] and sname_match(q,inv):
                inv["bank"]=q; q["sinv"]=inv; q["sinvtype"]="날짜+금액 일치"; break
pairs=[]
for inv in SINV:
    if inv["bank"] or inv["합계"]<=0: continue
    for q in QIx:
        if "sinv" in q or inv["합계"]!=q["입금"] or not sname_match(q,inv): continue
        pairs.append((abs((q["t"].date()-inv["d"]).days),id(inv),id(q),inv,q))
pairs.sort(key=lambda p:p[0])
for d,_,_,inv,q in pairs:
    if inv["bank"] or "sinv" in q: continue
    inv["bank"]=q; q["sinv"]=inv; q["sinvtype"]="금액+상호 일치"
for q in QIx:
    if "sinv" in q: continue
    for inv in SINV:
        if inv["bank"] is None and inv["합계"]>0 and inv["d"]==q["t"].date() and inv["합계"]==q["입금"]:
            inv["bank"]=q; q["sinv"]=inv; q["sinvtype"]="날짜+금액 일치"; break
# 분할입금: 같은 거래처 입금 2~10건 합 = 계산서
for inv in SINV:
    if inv["bank"] or inv["합계"]<=0: continue
    cands=[q for q in QIx if "sinv" not in q and sname_match(q,inv) and -180<=(q["t"].date()-inv["d"]).days<=240 and q["입금"]<inv["합계"]]
    cands.sort(key=lambda q:abs((q["t"].date()-inv["d"]).days)); cands=cands[:20]
    found=subset_sum(inv["합계"],cands,lambda q:q["입금"])
    if found and len(found)>=2:
        inv["bank"]=found[0]; inv["split"]=found
        for q in found: q["sinv"]=inv; q["sinvtype"]="분할 일치"
# 합산입금: 입금 1건 = 계산서 여러 장
for q in QIx:
    if "sinv" in q: continue
    cands=[inv for inv in SINV if not inv["bank"] and inv["합계"]>0 and sname_match(q,inv) and -240<=(inv["d"]-q["t"].date()).days<=180 and inv["합계"]<q["입금"]]
    cands.sort(key=lambda inv:abs((inv["d"]-q["t"].date()).days)); cands=cands[:20]
    found=subset_sum(q["입금"],cands,lambda i:i["합계"])
    if found and len(found)>=2:
        q["sinv"]=found[0]; q["sinvtype"]="합산 일치"; q["sinvs"]=found
        for inv in found: inv["bank"]=q; inv["merged"]=found
# 거래처별 합계 대조 (유효 계산서 합 == 해당 거래처 입금 합 이면 미매칭 행도 '거래처 합계 일치')
VMERGE_S={"채정원":"박효민"}   # 같은 거래처로 묶을 대표자명
def vkey_s(inv): k=nm(inv["대표"]) or nm(inv["상호"]); return VMERGE_S.get(k,k)
VS={}
for inv in SINV:
    if inv.get("cancel") or inv["합계"]<=0: continue
    VS.setdefault(vkey_s(inv),[]).append(inv)
vendor_ok_s=0
for k,invs in VS.items():
    deps=[q for q in QIx if any(sname_match(q,i) for i in invs) and ("sinv" not in q or q["sinv"] in invs or (q.get("sinvs") and all(x in invs for x in q["sinvs"])))]
    if sum(i["합계"] for i in invs)==sum(q["입금"] for q in deps) and deps:
        for i in invs:
            if not i["bank"]: i["bank"]=deps[0]; i["vendor_ok"]=True
        for q in deps:
            if "sinv" not in q: q["sinv"]=invs[0]; q["sinvtype"]="거래처 합계 일치"
        vendor_ok_s+=1
print("매출 거래처합계 일치 거래처",vendor_ok_s)
print("매출계산서",len(SINV),"상계",sum(1 for i in SINV if i.get("cancel")),"매칭",sum(1 for i in SINV if i["bank"] and i["bank"]!="X"),"/ 유효양수",sum(1 for i in SINV if i["합계"]>0 and not i.get("cancel")))

# ---------- 워크북 ----------
wb=openpyxl.Workbook()
hdr_font=Font(bold=True,color="FFFFFF"); hdr_fill=PatternFill("solid",fgColor="1F4E78")
fill=PatternFill("solid",fgColor="E2EFDA")                                   # 입금
mfill=PatternFill("solid",fgColor="C6EFCE"); mfont=Font(color="006100")      # 계산서 일치
yfill=PatternFill("solid",fgColor="FFEB9C"); yfont=Font(color="9C5700")      # 분할
bfill=PatternFill("solid",fgColor="DDEBF7"); bfont=Font(color="1F4E78")      # 합산
pfill=PatternFill("solid",fgColor="FFC7CE"); pfont=Font(color="9C0006")      # 미정산
gfill=PatternFill("solid",fgColor="F2F2F2"); gfont=Font(color="A6A6A6")      # 은행비용
vfill=PatternFill("solid",fgColor="E4DFEC"); vfont=Font(color="5F497A")      # 거래처 합계 일치
PINK_VENDORS=["마감자재"]
def style(ws, widths, money_cols, freeze="A2"):
    for c in ws[1]:
        c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(horizontal="center",vertical="center")
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for col in money_cols:
        for c in ws[get_column_letter(col)][1:]: c.number_format='#,##0'
    ws.freeze_panes=freeze; ws.auto_filter.ref=ws.dimensions
def datefmt(ws,col):
    for c in ws[col][1:]:
        if not isinstance(c.value,str): c.number_format="yyyy-mm-dd"
def tcols(t,keys): return [t[k] if t else None for k in keys]

# 1) 통합내역
ws=wb.active; ws.title="통합내역"
ws.append(["No","거래일자","시간","구분","적요","입금액","출금액","잔액","내용(통장표시)","거래점명",
   "수취인성명","입금은행","입금계좌","이체금액","수수료","거래메모","이체구분","입금통장표시내용","이체매칭",
   "세금계산서 공급자","세금계산서 합계","세금계산서 작성일자","세금계산서 승인번호","계산서 매칭유형"])
for i,q in enumerate(QM,1):
    t=q.get("T"); inv=q.get("inv"); invs=q.get("invs")
    ws.append([i, q["t"].date(), q["t"].strftime("%H:%M:%S"), "입금" if q["입금"] else "출금", q["적요"], q["입금"] or None, q["출금"] or None, q["잔액"], q["내용"], q["점"],
               *tcols(t,["수취인","은행","계좌","금액","수수료","메모","구분","입금표시"]), "이체결과 연결" if t else None,
               (inv["상호"]+f" 외 {len(invs)-1}건" if invs else (inv["상호"] if inv else None)),
               (sum(i["합계"] for i in invs) if invs else (inv["합계"] if inv else None)),
               (", ".join(i["d"].strftime("%m/%d") for i in invs) if invs else (inv["d"] if inv else None)),
               (" / ".join(i["승인"] for i in invs) if invs else (inv["승인"] if inv else None)), q.get("invtype")])
datefmt(ws,"B"); datefmt(ws,"V")
style(ws,[6,11,9,6,8,13,13,14,22,9,22,10,17,13,8,24,14,16,12,24,13,11,28,14],[6,7,8,14,15,21],"C2")
for row in ws.iter_rows(min_row=2):
    mt=row[23].value
    if row[3].value=="입금":
        for c in row: c.fill=fill
    elif mt=="분할 일치":
        for c in row: c.fill=yfill; c.font=yfont
    elif mt=="합산 일치":
        for c in row: c.fill=bfill; c.font=bfont
    elif mt=="거래처 합계 일치":
        row[20].value=None; row[21].value=None; row[22].value=None
        for c in row: c.fill=vfill; c.font=vfont
    elif str(row[8].value or "").upper().startswith("SMS"):
        row[23].value="은행비용(SMS)"
        for c in row: c.fill=gfill; c.font=gfont
    elif not row[19].value and any(v in str(row[8].value or "")+str(row[10].value or "") for v in PINK_VENDORS):
        row[23].value="미정산"
        for c in row: c.fill=pfill; c.font=pfont
    elif row[19].value:
        for c in row: c.fill=mfill; c.font=mfont

# 2) 입금내역
wd=wb.create_sheet("입금내역")
wd.append(["No","거래일자","시간","적요","입금액","잔액","입금자(내용)","거래점명","매출계산서 공급받는자","매출계산서 합계","매출계산서 작성일자","계산서 매칭유형"])
for i,q in enumerate(QI,1):
    inv=q.get("sinv"); invs=q.get("sinvs")
    wd.append([i,q["t"].date(),q["t"].strftime("%H:%M:%S"),q["적요"],q["입금"],q["잔액"],q["내용"],q["점"],
               ((inv["상호"] or inv["대표"])+f" 외 {len(invs)-1}건" if invs else ((inv["상호"] or inv["대표"]) if inv else None)),
               (sum(i["합계"] for i in invs) if invs else (inv["합계"] if inv else None)),
               (", ".join(i["d"].strftime("%m/%d") for i in invs) if invs else (inv["d"] if inv else None)), q.get("sinvtype")])
datefmt(wd,"B"); datefmt(wd,"K"); style(wd,[6,11,9,8,15,14,24,9,24,14,11,14],[5,6,10],"C2")
for row in wd.iter_rows(min_row=2):
    mt=row[11].value
    if mt=="분할 일치":
        for c in row: c.fill=yfill; c.font=yfont
    elif mt=="합산 일치":
        for c in row: c.fill=bfill; c.font=bfont
    elif mt=="거래처 합계 일치":
        for c in row: c.fill=vfill; c.font=vfont
    elif mt:
        for c in row: c.fill=mfill; c.font=mfont
    else:
        for c in row: c.fill=fill

# 매출세금계산서 탭
wsl=wb.create_sheet("매출세금계산서")
wsl.append(["No","작성일자","발급일자","공급받는자 등록번호","공급받는자 상호","대표자명","품목명","공급가액","세액","합계금액","영수/청구","승인번호","은행 매칭","은행 입금일자","은행 입금액","은행 입금자"])
for i,inv in enumerate(SINV,1):
    b=inv["bank"] if inv["bank"]!="X" else None
    if "split" in inv:
        mt="분할 일치"; bd=", ".join(q["t"].strftime("%Y-%m-%d") for q in inv["split"]); ba=sum(q["입금"] for q in inv["split"]); bc=" + ".join(f'{q["입금"]:,.0f}' for q in inv["split"])+f' ({b["내용"]})'
    elif "merged" in inv:
        mt="합산 일치"; bd=b["t"].date(); ba=b["입금"]; bc=f'{b["내용"]} (계산서 {len(inv["merged"])}장 합산 {sum(i["합계"] for i in inv["merged"]):,.0f})'
    else:
        mt=b.get("sinvtype") if b else None; bd=b["t"].date() if b else None; ba=b["입금"] if b else None; bc=b["내용"] if b else None
    if inv.get("cancel"): mt="취소/정정 상계"; bd=None; ba=None; bc=f'상계: {inv["cancel"]["d"]} {inv["cancel"]["합계"]:,.0f}'
    elif inv.get("vendor_ok"): mt="거래처 합계 일치"; bd=None; ba=None; bc="거래처 총액 기준 정산"
    elif inv["합계"]<0: mt="수정/취소(음수)"
    wsl.append([i,inv["d"],inv["발급"],inv["등록"],inv["상호"],inv["대표"],inv["품목"],inv["공급가액"],inv["세액"],inv["합계"],inv["영청"],inv["승인"],mt,bd,ba,bc])
datefmt(wsl,"B"); datefmt(wsl,"N")
style(wsl,[6,11,11,16,26,12,26,14,12,14,8,28,12,11,14,24],[8,9,10,15],"C2")
for row in wsl.iter_rows(min_row=2):
    mt=row[12].value
    if mt=="분할 일치":
        for c in row: c.fill=yfill; c.font=yfont
    elif mt=="합산 일치":
        for c in row: c.fill=bfill; c.font=bfont
    elif mt in ("수정/취소(음수)","취소/정정 상계"):
        for c in row: c.fill=gfill; c.font=gfont
    elif mt=="거래처 합계 일치":
        for c in row: c.fill=vfill; c.font=vfont
    elif mt:
        for c in row: c.fill=mfill; c.font=mfont
    else:
        for c in row: c.fill=pfill; c.font=pfont

# 3) 매입세금계산서
wi=wb.create_sheet("매입세금계산서")
wi.append(["No","작성일자","발급일자","공급자 상호","대표자명","사업자등록번호","품목명","공급가액","세액","합계금액","종류","구분","영수/청구","비고","승인번호","은행 매칭","은행 거래일자","은행 출금액","은행 내용"])
for i,inv in enumerate(INV,1):
    b=inv["bank"]
    if "split" in inv:
        mt="분할 일치"; bd=", ".join(q["t"].strftime("%Y-%m-%d") for q in inv["split"]); ba=sum(q["출금"] for q in inv["split"]); bc=" + ".join(f'{amt_of(q):,.0f}' for q in inv["split"])+f' ({b["내용"]})'
    elif "merged" in inv:
        mt="합산 일치"; bd=b["t"].date(); ba=b["출금"]; bc=f'{b["내용"]} (계산서 {len(inv["merged"])}장 합산 {sum(i["합계"] for i in inv["merged"]):,.0f})'
    elif inv.get("vendor_ok"):
        mt="거래처 합계 일치"; bd=None; ba=None; bc="거래처 총액 기준 정산"
    else:
        mt=b.get("invtype") if b else None; bd=b["t"].date() if b else None; ba=b["출금"] if b else None; bc=b["내용"] if b else None
    wi.append([i,inv["d"],inv["발급"],inv["상호"],inv["대표"],inv["사업자"],inv["품목"],inv["공급가액"],inv["세액"],inv["합계"],inv["종류"],inv["구분"],inv["영청"],inv["비고"],inv["승인"],mt,bd,ba,bc])
datefmt(wi,"B"); datefmt(wi,"Q")
style(wi,[6,11,11,26,10,14,30,13,11,13,10,6,8,16,28,10,11,13,24],[8,9,10,18],"C2")
for row in wi.iter_rows(min_row=2):
    mt=row[15].value
    if mt=="분할 일치":
        for c in row: c.fill=yfill; c.font=yfont
    elif mt=="합산 일치":
        for c in row: c.fill=bfill; c.font=bfont
    elif mt=="거래처 합계 일치":
        for c in row: c.fill=vfill; c.font=vfont
    elif not mt and any(v in str(row[3].value or "") for v in PINK_VENDORS):
        row[15].value="미정산"
        for c in row: c.fill=pfill; c.font=pfont
    elif mt:
        for c in row: c.fill=mfill; c.font=mfont

# 4) 세금·공과금
wt=wb.create_sheet("세금·공과금")
wt.append(["No","거래일자","시간","분류","적요","내용(통장표시)","입금액","출금액","잔액","수취인성명","입금은행","입금계좌","이체금액","수수료","거래메모","입금통장표시내용","세금계산서 공급자","세금계산서 합계","세금계산서 작성일자","계산서 매칭유형"])
for i,q in enumerate(QT,1):
    t=q.get("T"); inv=q.get("inv")
    wt.append([i,q["t"].date(),q["t"].strftime("%H:%M:%S"),q["tax"],q["적요"],q["내용"],q["입금"] or None,q["출금"] or None,q["잔액"],*tcols(t,["수취인","은행","계좌","금액","수수료","메모","입금표시"]),
               inv["상호"] if inv else None, inv["합계"] if inv else None, inv["d"] if inv else None, q.get("invtype")])
datefmt(wt,"B"); datefmt(wt,"S"); style(wt,[6,11,9,14,8,22,13,13,14,22,10,17,13,8,20,16,22,13,11,14],[7,8,9,13,14,18],"C2")
for row in wt.iter_rows(min_row=2):
    if row[19].value=="거래처 합계 일치":
        for c in row: c.fill=vfill; c.font=vfont
    elif row[19].value:
        for c in row: c.fill=mfill; c.font=mfont

# 5) 급여(김한승)
wp=wb.create_sheet("급여(김한승)")
wp.append(["No","거래일자","시간","적요","내용(통장표시)","출금액","이체금액","수수료","잔액","수취인성명","입금은행","입금계좌","거래메모","입금통장표시내용"])
for i,q in enumerate(QP,1):
    t=q.get("T")
    wp.append([i,q["t"].date(),q["t"].strftime("%H:%M:%S"),q["적요"],q["내용"],q["출금"],t["금액"] if t else None,t["수수료"] if t else None,q["잔액"],*tcols(t,["수취인","은행","계좌","메모","입금표시"])])
datefmt(wp,"B"); style(wp,[6,11,9,8,22,13,13,8,14,12,10,17,16,16],[6,7,8,9],"C2")

# 6) 급여명세(월별)
PAY=[]
for f in files:
    if "급여명세" not in N(f): continue
    wsP=openpyxl.load_workbook(f)["급여명세"]
    for r in wsP.iter_rows(min_row=2,values_only=True):
        if not isinstance(r[0],(int,float)): continue
        지급=sum(x or 0 for x in r[9:14]); 공제=sum(x or 0 for x in r[15:23])
        PAY.append(dict(귀속=r[1],y=r[2],m=r[3],지급일=r[4],성명=r[6],기본급=r[9],상여=r[10],식대=r[11],자가=r[12],육아=r[13],지급=지급,
                        국민=r[15],건강=r[16],고용=r[17],장기=r[18],소득=r[19],지방=r[20],연말소득=r[21],연말지방=r[22],공제=공제,차인=지급-공제,bank=None))
def pay_amt(q):
    t=q.get("T"); return t["금액"] if t else q["출금"]-500
def memo_txt(q):
    t=q.get("T"); return " ".join(str(x) for x in (q["내용"],t and t["메모"],t and t["입금표시"]) if x)
def memo_month(q):
    m=memo_txt(q).replace(" ","")
    r=re.search(r"(\d{2})(\d{2})월",m)
    if r: return 2000+int(r.group(1)),int(r.group(2))
    r=re.search(r"(\d{1,2})월",m) or re.search(r"급여(\d{1,2})",m) or re.search(r"(\d{1,2})급여",m)
    if r: return None,int(r.group(1))
    return None,None
TOL=1000; used=set()
QPs=sorted(QP,key=lambda q:q["t"])
for q in sorted(QPs,key=lambda q:(memo_month(q)[0] is None, q["t"])):
    y,m=memo_month(q)
    if not m: continue
    c=[p for p in PAY if p["bank"] is None and p["m"]==m and (y is None or p["y"]==y) and abs(pay_amt(q)-p["차인"])<=TOL and 0<=(q["t"].date()-dt.date(p["y"],p["m"],1)).days<=400]
    if c:
        p=min(c,key=lambda p:(q["t"].date()-dt.date(p["y"],p["m"],1)).days); p["bank"]=q; used.add(id(q))
for q in QPs:
    if id(q) in used: continue
    c=[p for p in PAY if p["bank"] is None and abs(pay_amt(q)-p["차인"])<=TOL and dt.date(p["y"],p["m"],1)<=q["t"].date()]
    if c:
        p=min(c,key=lambda p:(p["y"],p["m"])); p["bank"]=q; used.add(id(q))
wm=wb.create_sheet("급여명세(월별)")
wm.append(["귀속월","지급일(명세)","성명","기본급","상여","식대","자가운전보조금","육아수당","지급합계","국민연금","건강보험","고용보험","장기요양","소득세","지방소득세","연말정산소득세","연말정산지방소득세","공제합계","차인지급액","은행 이체일","은행 이체원금","차이","이체메모","지급상태"])
for p in PAY:
    b=p["bank"]
    wm.append([p["귀속"],p["지급일"],p["성명"],p["기본급"],p["상여"],p["식대"],p["자가"],p["육아"],p["지급"],p["국민"],p["건강"],p["고용"],p["장기"],p["소득"],p["지방"],p["연말소득"],p["연말지방"],p["공제"],p["차인"],
               b["t"].date() if b else None, pay_amt(b) if b else None, (pay_amt(b)-p["차인"]) if b else None, memo_txt(b) if b else None, "지급확인" if b else "미확인"])
datefmt(wm,"T")
style(wm,[9,11,8,11,8,9,12,9,11,10,10,8,9,9,10,12,14,11,12,11,12,8,22,9],list(range(4,20))+[21,22],"D2")
for row in wm.iter_rows(min_row=2):
    if row[23].value=="지급확인":
        for c in row: c.fill=mfill; c.font=mfont
    else:
        for c in row: c.fill=pfill; c.font=pfont
extra=[q for q in QP if id(q) not in used]
if extra:
    wm.append([]); wm.append(["명세와 연결 안 된 은행 급여이체"]); wm[wm.max_row][0].font=Font(bold=True)
    for q in extra: wm.append([None]*19+[q["t"].date(),pay_amt(q),None,memo_txt(q),"명세 없음"])

# 7) 대여금(김한승)
wl=wb.create_sheet("대여금(김한승)")
wl.append(["No","거래일자","시간","연도","구분","적요","내용(통장표시)","대표→회사 입금","회사→대표 출금(원금)","수수료","거래메모","누적잔액(대표 순입금)"])
bal=0; ysum={}
def loan_split(q):
    t=q.get("T"); fee=(t["수수료"] if t else 0) or 0
    out=(t["금액"] if t else q["출금"]) if q["출금"] else 0
    if q["출금"] and not t and out%1000==500: out-=500; fee=500
    return out,fee
for i,q in enumerate(QL,1):
    t=q.get("T"); out,fee=loan_split(q); bal+=q["입금"]-out
    y=q["t"].year; d=ysum.setdefault(y,[0,0,0]); d[0]+=q["입금"]; d[1]+=out; d[2]+=fee
    wl.append([i,q["t"].date(),q["t"].strftime("%H:%M:%S"),y,"입금" if q["입금"] else "출금",q["적요"],q["내용"],q["입금"] or None,out or None,fee or None,t["메모"] if t else None,bal])
datefmt(wl,"B"); style(wl,[6,11,9,6,6,8,20,15,17,8,16,18],[8,9,10,12],"C2")
for row in wl.iter_rows(min_row=2):
    if row[4].value=="입금":
        for c in row: c.fill=fill
wl.append([]); wl.append(["연도별 요약",None,None,"연도",None,None,None,"대표→회사 입금","회사→대표 출금(원금)","수수료",None,"연말 누적잔액"])
for c in wl[wl.max_row]: c.font=Font(bold=True)
cum=0
for y,d in sorted(ysum.items()):
    cum+=d[0]-d[1]; wl.append([None,None,None,y,None,None,None,d[0],d[1],d[2],None,cum])
    for col in (8,9,10,12): wl.cell(wl.max_row,col).number_format='#,##0'

# 8) 가수금 정산(2025~)
wg=wb.create_sheet("가수금 정산(2025~)")
wg.append(["No","거래일자","연도","구분","내용(통장표시)","거래메모","대표→회사 입금(가수금 증가)","회사→대표 출금(가수금 감소, 원금)","수수료","가수금 잔액"])
OPEN=42575546
wg.append([None,dt.date(2024,12,31),2024,"기초","2024.12.31 기준 가수금 (회계사 확인)",None,None,None,None,OPEN])
for c in wg[2]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="FFF2CC")
gbal=OPEN; gy={}; n=0
for q in QL:
    if q["t"].year<2025: continue
    t=q.get("T"); out,fee=loan_split(q); gbal+=q["입금"]-out; n+=1
    d=gy.setdefault(q["t"].year,[0,0,0,0]); d[0]+=q["입금"]; d[1]+=out; d[2]+=fee; d[3]=gbal
    wg.append([n,q["t"].date(),q["t"].year,"입금" if q["입금"] else "출금",q["내용"],t["메모"] if t else None,q["입금"] or None,out or None,fee or None,gbal])
datefmt(wg,"B"); style(wg,[6,11,6,6,20,16,22,26,8,16],[7,8,9,10],"C3")
for row in wg.iter_rows(min_row=3):
    if row[3].value=="입금":
        for c in row: c.fill=fill
wg.append([]); wg.append(["연도별 정산",None,"연도",None,None,None,"가수금 증가(입금)","가수금 감소(출금 원금)","수수료","연말 가수금 잔액"])
for c in wg[wg.max_row]: c.font=Font(bold=True)
wg.append([None,None,2024,None,None,None,None,None,None,OPEN])
for y,d in sorted(gy.items()): wg.append([None,None,y,None,None,None,d[0],d[1],d[2],d[3]])
for r in range(wg.max_row-len(gy),wg.max_row+1):
    for col in (7,8,9,10): wg.cell(r,col).number_format='#,##0'

# 9) 체크카드
wc=wb.create_sheet("체크카드")
wc.append(["No","거래일자","시간","구분","가맹점(내용)","사용금액","입금(취소/환불)","잔액","거래점명"])
for i,q in enumerate(QC,1):
    wc.append([i,q["t"].date(),q["t"].strftime("%H:%M:%S"),"입금" if q["입금"] else "출금",q["내용"],q["출금"] or None,q["입금"] or None,q["잔액"],q["점"]])
datefmt(wc,"B"); style(wc,[6,11,9,6,26,13,13,14,9],[6,7,8],"C2")
for row in wc.iter_rows(min_row=2):
    if row[3].value=="입금":
        for c in row: c.fill=fill

# 10) 원본통합
ws2=wb.create_sheet("이체결과(원본통합)")
ws2.append(["No","거래일시","이체구분","처리결과","입금은행","입금계좌","수취인성명","처리금액","수수료","거래메모","출금통장표시내용","입금통장표시내용","조회매칭"])
for i,t in enumerate(T,1):
    ws2.append([i,t["t"],t["구분"],t["결과"],t["은행"],t["계좌"],t["수취인"],t["금액"],t["수수료"],t["메모"],t["출금표시"],t["입금표시"],"O" if t["used"] else "미매칭"])
for c in ws2["B"][1:]: c.number_format="yyyy-mm-dd hh:mm:ss"
style(ws2,[6,19,15,8,10,17,22,13,8,24,16,16,9],[8,9])
ws3=wb.create_sheet("조회(원본통합)")
ws3.append(["No","거래일시","적요","입금액","출금액","내용","잔액","거래점명"])
for i,q in enumerate(Q,1):
    ws3.append([i,q["t"],q["적요"],q["입금"] or None,q["출금"] or None,q["내용"],q["잔액"],q["점"]])
for c in ws3["B"][1:]: c.number_format="yyyy-mm-dd hh:mm:ss"
style(ws3,[6,19,8,13,13,22,14,9],[4,5,7])

# 11) 월별요약
ws4=wb.create_sheet("월별요약")
ws4.append(["연월","입금건수","입금합계","출금건수","출금합계","순증감","월말잔액"])
m=OrderedDict()
for q in Q:
    k=q["t"].strftime("%Y-%m"); d=m.setdefault(k,[0,0,0,0,None])
    if q["입금"]: d[0]+=1; d[1]+=q["입금"]
    if q["출금"]: d[2]+=1; d[3]+=q["출금"]
    d[4]=q["잔액"]
for k,d in m.items(): ws4.append([k,d[0],d[1],d[2],d[3],d[1]-d[3],d[4]])
style(ws4,[9,9,15,9,15,15,15],[3,5,6,7])

# ---------- 최신 날짜가 위로 ----------
def reverse_rows(ws,r1,r2):
    rows=[]
    for r in range(r1,r2+1):
        rows.append([(c.value,copy(c.number_format),copy(c.fill),copy(c.font),copy(c.alignment)) for c in ws[r]])
    rows.reverse()
    for i,row in enumerate(rows):
        for j,(v,nf,fl,fo,al) in enumerate(row):
            c=ws.cell(r1+i,j+1); c.value=v; c.number_format=nf; c.fill=fl; c.font=fo; c.alignment=al
def data_end(ws,start):
    r=start
    while r<=ws.max_row and any(c.value is not None for c in ws[r]): r+=1
    return r-1
for name in wb.sheetnames:
    w=wb[name]
    start=3 if name.startswith("가수금") else 2
    end=data_end(w,start)
    if end>start: reverse_rows(w,start,end)
    if w.cell(1,1).value=="No":
        for k,r in enumerate(range(start,end+1),1): w.cell(r,1).value=k

wb.move_sheet("매출세금계산서", offset=-(len(wb.sheetnames)-1-wb.sheetnames.index("입금내역")-1))
