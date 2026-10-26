"""은행에서 내려받은 거래내역 파일(xlsx / xls / HTML로 된 xls / csv) → 공통 거래 행.

머리글 행을 글자로 찾아 열을 맞추므로 은행 양식이 조금 바뀌어도 대부분 읽힌다.
돌려주는 행은 report.py 의 CODEF 결과와 같은 모양이다.
"""

import csv
import re
from datetime import date
from pathlib import Path

COLUMN_KEYS = {
    # 앞에 있는 후보일수록 우선
    "date": ["거래일시", "거래일자", "거래일", "일자", "거래날짜", "날짜"],
    "time": ["거래시간", "거래시각", "시간", "시각"],
    "out": ["출금금액", "출금액", "출금", "찾으신금액", "지급금액", "지급"],
    "in": ["입금금액", "입금액", "입금", "맡기신금액"],
    "balance": ["거래후잔액", "거래후 잔액", "잔액"],
    "desc1": ["받는분", "보낸분", "기재내용", "거래내용", "내용", "적요"],
    "desc2": ["거래구분", "구분", "거래방법", "거래점", "취급점"],
    "desc3": ["메모", "통장메모", "비고", "받는분통장표시", "의뢰인", "적요"],
}
ACCOUNT_RE = re.compile(r"\d{3}-\d{2,3}-\d{5,7}|\d{3}-\d{6}-\d{2,5}")


class ParseError(Exception):
    pass


# ------------------------------------------------------------------ 파일 → 표


def read_table(path: Path) -> list[list[str]]:
    data = path.read_bytes()
    head = data[:2048].lstrip().lower()
    if data[:4] == b"PK\x03\x04":
        return _read_xlsx(path)
    if data[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
        return _read_xls(path)
    if head.startswith((b"<", b"\xef\xbb\xbf<")) or b"<table" in head or b"<html" in head:
        return _read_html(data)
    return _read_csv(data)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M:%S") if hasattr(value, "hour") else value.strftime("%Y-%m-%d")
    return str(value).strip()


def _read_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        rows += [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    return rows


def _read_xls(path: Path) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    rows = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    row.append(_cell(cell.value))
            rows.append(row)
    return rows


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_html(data: bytes) -> list[list[str]]:
    from lxml import html

    text = _decode(data)
    tree = html.fromstring(text)
    rows = []
    for tr in tree.iter("tr"):
        rows.append([" ".join(td.text_content().split()) for td in tr if td.tag in ("td", "th")])
    return rows


def _read_csv(data: bytes) -> list[list[str]]:
    return [[c.strip() for c in row] for row in csv.reader(_decode(data).splitlines())]


# ------------------------------------------------------------------ 표 → 거래


def _norm(text: str) -> str:
    return re.sub(r"[\s()\[\]:·]", "", text)


def find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:60]):
        names = [_norm(c) for c in row]
        mapping: dict[str, int] = {}
        for key, candidates in COLUMN_KEYS.items():
            for candidate in candidates:
                hit = next((i for i, n in enumerate(names) if n and _norm(candidate) in n and i not in mapping.values()), None)
                if hit is not None:
                    mapping[key] = hit
                    break
        if "date" in mapping and ("out" in mapping or "in" in mapping):
            return index, mapping
    raise ParseError("거래일/입금/출금 머리글을 찾지 못했습니다.")


def _amount(text: str) -> int:
    digits = re.sub(r"[^\d-]", "", text or "")
    try:
        return int(digits) if digits not in ("", "-") else 0
    except ValueError:
        return 0


def _date_time(date_text: str, time_text: str) -> tuple[str, str]:
    numbers = re.findall(r"\d+", date_text or "")
    if not numbers:
        return "", ""
    joined = "".join(numbers)
    if len(numbers) >= 3 and len(numbers[0]) == 4:
        ymd = f"{numbers[0]}{int(numbers[1]):02d}{int(numbers[2]):02d}"
        rest = numbers[3:]
    elif len(joined) >= 8:
        ymd, rest = joined[:8], [joined[8:]] if len(joined) > 8 else []
    else:
        return "", ""
    clock = "".join(re.findall(r"\d+", time_text or "")) or "".join(f"{int(n):02d}" for n in rest)
    return ymd, (clock + "000000")[:6] if clock else ""


def parse_file(path: Path, start: date | None = None, end: date | None = None) -> list[dict]:
    rows = read_table(path)
    header, cols = find_header(rows)
    account = ""
    for row in rows[:header]:
        match = ACCOUNT_RE.search(" ".join(row))
        if match:
            account = match.group(0)
            break
    account = account or path.stem

    def get(row, key):
        index = cols.get(key)
        return row[index] if index is not None and index < len(row) else ""

    result = []
    for row in rows[header + 1:]:
        ymd, hms = _date_time(get(row, "date"), get(row, "time"))
        if not ymd:
            continue  # 합계·빈 줄
        if start and end and not (start.strftime("%Y%m%d") <= ymd <= end.strftime("%Y%m%d")):
            continue
        result.append(
            {
                "account": account,
                "account_name": "",
                "date": ymd,
                "time": hms,
                "out": _amount(get(row, "out")),
                "in": _amount(get(row, "in")),
                "balance": _amount(get(row, "balance")),
                "desc1": get(row, "desc1"),
                "desc2": get(row, "desc2"),
                "desc3": get(row, "desc3"),
                "desc4": "",
            }
        )
    return result


def parse_files(paths, start: date | None = None, end: date | None = None) -> tuple[list[dict], list[str]]:
    rows, warnings = [], []
    for path in paths:
        try:
            rows += parse_file(Path(path), start, end)
        except Exception as error:
            warnings.append(f"은행 파일 {Path(path).name} 을(를) 읽지 못했습니다: {error}")
    return rows, warnings


# ------------------------------------------------------------------ 이체결과 등 일반 표

TABLE_KEYS = ["이체일", "이체일자", "이체일시", "처리일시", "거래일시", "신청일", "이체금액", "금액",
              "받는분", "수취인", "입금계좌", "입금은행", "처리결과", "결과", "수수료", "출금계좌", "상태"]


def parse_table(path: Path, start: date | None = None, end: date | None = None) -> tuple[list[str], list[list]]:
    """머리글 행을 찾아 표를 그대로 옮긴다(이체결과조회처럼 양식을 모르는 파일용).

    금액 열은 숫자로 바꾸고, 날짜 열이 있으면 기간 밖 행을 뺀다."""
    rows = read_table(path)
    header_index = None
    for index, row in enumerate(rows[:60]):
        hits = sum(1 for cell in row if cell and any(_norm(k) in _norm(cell) for k in TABLE_KEYS))
        if hits >= 3:
            header_index = index
            break
    if header_index is None:
        raise ParseError("이체결과 표의 머리글(이체일·금액·받는분 등)을 찾지 못했습니다.")
    headers = [c or f"열{i + 1}" for i, c in enumerate(rows[header_index])]
    while headers and headers[-1].startswith("열"):
        headers.pop()
    width = len(headers)
    date_col = next((i for i, h in enumerate(headers) if re.search(r"일자|일시|이체일|처리일|신청일", h)), None)
    money_cols = {i for i, h in enumerate(headers) if re.search(r"금액|수수료|잔액", h)}
    body = []
    for row in rows[header_index + 1:]:
        cells = (row + [""] * width)[:width]
        if sum(1 for c in cells if c) < 2 or _norm(cells[0]).startswith(("합계", "총계", "소계")):
            continue
        if date_col is not None and start and end:
            ymd, _ = _date_time(cells[date_col], "")
            if ymd and not (start.strftime("%Y%m%d") <= ymd <= end.strftime("%Y%m%d")):
                continue
        body.append([_amount(c) if i in money_cols and re.fullmatch(r"[\d,\s-]+원?", c or "") else c
                     for i, c in enumerate(cells)])  # fmt: skip
    return headers, body


def parse_tables(paths, start: date | None = None, end: date | None = None):
    """여러 파일을 한 표로. 머리글이 다르면 첫 파일 기준. (headers, rows, warnings)"""
    headers, rows, warnings = None, [], []
    for path in paths:
        try:
            h, body = parse_table(Path(path), start, end)
        except Exception as error:
            warnings.append(f"이체결과 파일 {Path(path).name} 을(를) 읽지 못했습니다: {error}")
            continue
        headers = headers or h
        rows += body
    return headers, rows, warnings


def classify(paths) -> tuple[list[Path], list[Path]]:
    """직접 고른 파일을 거래내역 / 이체결과로 나눈다."""
    transactions, others = [], []
    for path in map(Path, paths):
        try:
            find_header(read_table(path))
            transactions.append(path)
        except Exception:
            others.append(path)
    return transactions, others
