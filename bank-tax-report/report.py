"""신한은행 기업 계좌 거래·이체 내역 + 홈택스 세금계산서 발급/수취 내역 → 엑셀 한 파일.

같은 법인 공동인증서(NPKI: signCert.der / signPri.key) 하나로 두 곳을 조회한다.

  은행(기본, browser): 앱 전용 크롬 창으로 신한 기업뱅킹을 열고, 사람이 로그인하면
        거래내역조회 → 엑셀저장을 눌러 받은 파일을 읽는다(shinhan_browser.py, bank_excel.py).
  은행(codef): CODEF API(https://codef.io)가 공동인증서로 신한은행 기업뱅킹에 로그인해
        법인 수시입출 계좌의 거래내역을 돌려준다.
  은행(files): 직접 내려받아 둔 거래내역 엑셀 파일을 읽는다.
  이체(출금) 내역은 거래내역에서 출금 건을 뽑는다.
  홈택스: 오픈소스 `hometax` CLI(https://github.com/hjsh200219/hometax)가 공동인증서로
        홈택스에 로그인해 전자세금계산서 매출(발급)·매입(수취) 목록을 돌려준다.

환경 변수 (.env 파일 또는 시스템 환경 변수):
  CODEF_CLIENT_ID, CODEF_CLIENT_SECRET, CODEF_PUBLIC_KEY   CODEF 키 관리 화면의 값
  CODEF_ENV            demo(기본, 개발용 무료) / production(정식)
  CERT_DIR             인증서 폴더(signCert.der, signPri.key가 들어 있는 폴더)
  CERT_PASSWORD        인증서 비밀번호(비워 두면 실행할 때 물어봄)
  SHINHAN_ACCOUNTS     조회할 계좌번호(쉼표 구분, '-' 없이). 비우면 보유 계좌 전체
  OUTPUT_DIR           엑셀 저장 폴더(기본: 이 스크립트 옆 output/)

사용:
  python report.py                 # 오늘 포함 최근 7일
  python report.py --days 14
  python report.py --from 2026-10-01 --to 2026-10-07
  python report.py --skip-bank     # 홈택스만 / --skip-hometax 은행만
  python report.py --bank codef    # 은행을 CODEF 로
  python report.py --bank files --bank-file 거래내역.xls
  python report.py --demo          # 가짜 데이터로 엑셀 모양만 확인(네트워크·인증서 불필요)
"""

import argparse
import base64
import getpass
import json
import os
import shutil
import subprocess
import sys
import urllib.parse
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
KST = timezone(timedelta(hours=9))

SHINHAN = "0088"
CODEF_OAUTH = "https://oauth.codef.io/oauth/token"
CODEF_DOMAINS = {
    "demo": "https://development.codef.io",
    "production": "https://api.codef.io",
    "sandbox": "https://sandbox.codef.io",
}
PATH_ACCOUNT_CREATE = "/v1/account/create"
PATH_ACCOUNT_ADD = "/v1/account/add"
PATH_ACCOUNT_LIST = "/v1/kr/bank/b/account/account-list"
PATH_TRANSACTIONS = "/v1/kr/bank/b/account/transaction-list"


class AppError(Exception):
    pass


# ------------------------------------------------------------------ 설정


def load_dotenv(path: Path) -> None:
    """KEY=VALUE 형식의 .env를 읽어 비어 있는 환경 변수만 채운다."""
    if not path.is_file():
        return
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip('"').strip("'")
        if value and not os.environ.get(key.strip()):
            os.environ[key.strip()] = value


def env(name: str, default: str = "") -> str:
    return os.environ.get(name, default).strip()


def today_kst() -> date:
    return datetime.now(KST).date()


def parse_period(args) -> tuple[date, date]:
    if args.start or args.end:
        if not (args.start and args.end):
            raise AppError("--from 과 --to 는 함께 지정하세요.")
        start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    else:
        end = today_kst()
        start = end - timedelta(days=args.days - 1)
    if start > end:
        raise AppError("시작일이 종료일보다 늦습니다.")
    if end > today_kst():
        raise AppError("미래 날짜는 조회할 수 없습니다.")
    return start, end


def cert_files(cert_dir: Path) -> tuple[Path, Path]:
    der, key = cert_dir / "signCert.der", cert_dir / "signPri.key"
    if not der.is_file() or not key.is_file():
        raise AppError(
            f"인증서 폴더에 signCert.der / signPri.key 가 없습니다: {cert_dir}\n"
            "  예) C:\\Users\\<사용자>\\AppData\\LocalLow\\NPKI\\yessign\\USER\\cn=...(법인명)..."
        )
    return der, key


def data_dir() -> Path:
    """설정·상태 파일 위치. 맥 앱은 BTR_HOME 으로 ~/Library/Application Support 아래를 쓴다."""
    return Path(os.environ.get("BTR_HOME") or HERE)


def state_file() -> Path:
    return data_dir() / ".state.json"  # CODEF connectedId 보관(비밀번호·인증서는 저장하지 않음)


def load_state() -> dict:
    try:
        return json.loads(state_file().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_state(state: dict) -> None:
    state_file().parent.mkdir(parents=True, exist_ok=True)
    state_file().write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


# ------------------------------------------------------------------ 신한은행 (CODEF)


class Codef:
    def __init__(self, client_id: str, client_secret: str, public_key: str, mode: str):
        import requests  # 데모 모드에선 필요 없도록 지연 import

        if not (client_id and client_secret and public_key):
            raise AppError("CODEF_CLIENT_ID / CODEF_CLIENT_SECRET / CODEF_PUBLIC_KEY 를 설정하세요.")
        if mode not in CODEF_DOMAINS:
            raise AppError(f"CODEF_ENV 는 {', '.join(CODEF_DOMAINS)} 중 하나여야 합니다.")
        self.http = requests.Session()
        self.client_id, self.client_secret = client_id, client_secret
        self.public_key = public_key
        self.base = CODEF_DOMAINS[mode]
        self.token = ""

    def encrypt(self, text: str) -> str:
        """CODEF 공개키(RSA, PKCS#1 v1.5)로 비밀번호를 암호화한다."""
        from cryptography.hazmat.primitives.asymmetric import padding
        from cryptography.hazmat.primitives.serialization import load_der_public_key

        key = load_der_public_key(base64.b64decode(self.public_key))
        return base64.b64encode(key.encrypt(text.encode("utf-8"), padding.PKCS1v15())).decode()

    def _issue_token(self) -> None:
        basic = base64.b64encode(f"{self.client_id}:{self.client_secret}".encode()).decode()
        response = self.http.post(
            CODEF_OAUTH,
            headers={
                "Content-Type": "application/x-www-form-urlencoded",
                "Authorization": "Basic " + basic,
            },
            data="grant_type=client_credentials&scope=read",
            timeout=30,
        )
        if response.status_code != 200:
            raise AppError(f"CODEF 토큰 발급 실패({response.status_code}): 키를 확인하세요.")
        self.token = response.json()["access_token"]

    def call(self, path: str, body: dict) -> dict:
        """CODEF는 요청·응답 본문을 URL 인코딩한 JSON으로 주고받는다."""
        if not self.token:
            self._issue_token()
        for attempt in range(2):
            response = self.http.post(
                self.base + path,
                headers={"Content-Type": "application/json", "Authorization": "Bearer " + self.token},
                data=urllib.parse.quote(json.dumps(body, ensure_ascii=False)),
                timeout=120,
            )
            if response.status_code == 401 and attempt == 0:
                self._issue_token()
                continue
            break
        if response.status_code != 200:
            raise AppError(f"CODEF 호출 실패 {path} ({response.status_code})")
        return json.loads(urllib.parse.unquote_plus(response.text))


def codef_error(payload: dict) -> str:
    result = payload.get("result", {})
    message = f"{result.get('code')} {result.get('message', '')} {result.get('extraMessage', '')}"
    data = payload.get("data")
    for item in data.get("errorList", []) if isinstance(data, dict) else []:
        message += f"\n  - {item.get('code')} {item.get('message')} {item.get('extraMessage', '')}"
    return message.strip()


def ensure_connected_id(codef: Codef, der: Path, key: Path, password: str, mode: str) -> str:
    """신한은행 기업 계정(공동인증서)을 CODEF에 등록하고 connectedId를 받는다. 한 번만 하면 된다."""
    state = load_state()
    saved = state.get(f"connected_id_{mode}")
    if saved and state.get(f"shinhan_registered_{mode}"):
        return saved
    account = {
        "countryCode": "KR",
        "businessType": "BK",
        "clientType": "B",  # 기업(법인·개인사업자)
        "organization": SHINHAN,
        "loginType": "0",  # 0: 공동인증서
        "password": codef.encrypt(password),
        "derFile": base64.b64encode(der.read_bytes()).decode(),
        "keyFile": base64.b64encode(key.read_bytes()).decode(),
    }
    if saved:
        payload = codef.call(PATH_ACCOUNT_ADD, {"connectedId": saved, "accountList": [account]})
    else:
        payload = codef.call(PATH_ACCOUNT_CREATE, {"accountList": [account]})
    data = payload.get("data") or {}
    connected_id = data.get("connectedId") or saved
    succeeded = any(item.get("organization") == SHINHAN for item in data.get("successList", []))
    if not connected_id or not succeeded:
        raise AppError("신한은행 공동인증서 등록 실패:\n" + codef_error(payload))
    state[f"connected_id_{mode}"] = connected_id
    state[f"shinhan_registered_{mode}"] = True
    save_state(state)
    return connected_id


def list_accounts(codef: Codef, connected_id: str) -> list[str]:
    payload = codef.call(PATH_ACCOUNT_LIST, {"connectedId": connected_id, "organization": SHINHAN})
    if payload.get("result", {}).get("code") != "CF-00000":
        raise AppError("신한은행 계좌 목록 조회 실패:\n" + codef_error(payload))
    deposits = (payload.get("data") or {}).get("resDepositTrust", [])
    return [item["resAccount"] for item in deposits if item.get("resAccount")]


def to_int(value) -> int:
    try:
        return int(str(value or "0").replace(",", "").strip() or 0)
    except ValueError:
        return 0


def fetch_bank(start: date, end: date, cert_dir: Path, password: str) -> tuple[list[dict], list[str]]:
    mode = env("CODEF_ENV", "demo")
    codef = Codef(env("CODEF_CLIENT_ID"), env("CODEF_CLIENT_SECRET"), env("CODEF_PUBLIC_KEY"), mode)
    der, key = cert_files(cert_dir)
    connected_id = ensure_connected_id(codef, der, key, password, mode)

    accounts = [a.replace("-", "").strip() for a in env("SHINHAN_ACCOUNTS").split(",") if a.strip()]
    if not accounts:
        accounts = list_accounts(codef, connected_id)
        print(f"  보유 계좌 {len(accounts)}개: {', '.join(accounts)}")

    rows, warnings = [], []
    for account in accounts:
        payload = codef.call(
            PATH_TRANSACTIONS,
            {
                "connectedId": connected_id,
                "organization": SHINHAN,
                "account": account,
                "startDate": start.strftime("%Y%m%d"),
                "endDate": end.strftime("%Y%m%d"),
                "orderBy": "0",  # 0: 최신순
            },
        )
        if payload.get("result", {}).get("code") != "CF-00000":
            warnings.append(f"계좌 {account} 조회 실패: {codef_error(payload)}")
            continue
        data = payload.get("data") or {}
        for tr in data.get("resTrHistoryList", []):
            rows.append(
                {
                    "account": data.get("resAccount") or account,
                    "account_name": data.get("resAccountName", ""),
                    "date": tr.get("resAccountTrDate", ""),
                    "time": tr.get("resAccountTrTime", ""),
                    "out": to_int(tr.get("resAccountOut")),
                    "in": to_int(tr.get("resAccountIn")),
                    "balance": to_int(tr.get("resAfterTranBalance")),
                    "desc1": tr.get("resAccountDesc1", ""),
                    "desc2": tr.get("resAccountDesc2", ""),
                    "desc3": tr.get("resAccountDesc3", ""),
                    "desc4": tr.get("resAccountDesc4", ""),
                }
            )
    return rows, warnings


# ------------------------------------------------------------------ 홈택스 (hometax CLI)


def hometax_command() -> list[str]:
    # 맥 앱은 PATH가 짧으므로 지금 쓰는 파이썬(가상환경)의 bin 을 먼저 본다
    beside = Path(sys.executable).parent / ("hometax.exe" if os.name == "nt" else "hometax")
    found = str(beside) if beside.is_file() else shutil.which("hometax")
    if found:
        return [found]
    # pip/uv로 설치했지만 PATH에 없을 때
    return [sys.executable, "-m", "hometax_login.cli"]


def fetch_hometax(start: date, end: date, cert_dir: Path, password: str) -> dict[str, list[dict]]:
    cert_files(cert_dir)
    child_env = dict(os.environ, HOMETAX_PW=password, HOMETAX_NPKI_PATH=str(cert_dir))
    result = {}
    for direction in ("sales", "purchases"):
        command = [
            *hometax_command(), "invoices", "--json",
            "--from", start.isoformat(), "--to", end.isoformat(),
            "--direction", direction, "--basis", env("HOMETAX_DATE_BASIS", "issued"),
            "--cert", str(cert_dir), "--no-remember",
        ]  # fmt: skip
        try:
            done = subprocess.run(
                command, env=child_env, capture_output=True, text=True,
                encoding="utf-8", timeout=300, stdin=subprocess.DEVNULL,
            )  # fmt: skip
        except FileNotFoundError:
            raise AppError("hometax 도구가 설치되지 않았습니다. README의 설치 단계를 따라 주세요.") from None
        if done.returncode != 0:
            raise AppError(f"홈택스 조회 실패({direction}): {done.stderr.strip() or done.stdout.strip()}")
        result[direction] = json.loads(done.stdout)["items"]
    return result


# ------------------------------------------------------------------ 엑셀


def ymd(value: str) -> str:
    value = str(value or "")
    return f"{value[:4]}-{value[4:6]}-{value[6:8]}" if len(value) == 8 and value.isdigit() else value


def hms(value: str) -> str:
    value = str(value or "")
    return f"{value[:2]}:{value[2:4]}:{value[4:6]}" if len(value) == 6 and value.isdigit() else value


@dataclass
class Sheet:
    title: str
    headers: list[str]
    rows: list[list]
    money_cols: tuple[int, ...]  # 1부터


def build_workbook(path: Path, start: date, end: date, bank: list[dict] | None,
                   invoices: dict[str, list[dict]] | None, warnings: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bank_headers = ["계좌번호", "거래일", "시각", "출금액", "입금액", "거래후잔액", "적요/받는분", "거래구분", "메모", "기타"]

    def bank_row(r):
        return [r["account"], ymd(r["date"]), hms(r["time"]), r["out"], r["in"], r["balance"],
                r["desc1"], r["desc2"], r["desc3"], r["desc4"]]  # fmt: skip

    def invoice_rows(items):
        return [
            [i["issued_date"], i["written_date"], i.get("transmitted_date") or "",
             i.get("counterparty_name") or "", i.get("item_name") or "",
             i["supply_amount"], i["tax_amount"], i["total_amount"], i["approval_number"]]
            for i in items
        ]  # fmt: skip

    invoice_headers = lambda who: ["발급일", "작성일", "전송일", who, "품목", "공급가액", "세액", "합계금액", "승인번호"]  # noqa: E731

    sheets: list[Sheet] = []
    summary: list[list] = [["조회 기간", f"{start} ~ {end}", "", ""], ["", "", "", ""],
                           ["구분", "건수", "금액1", "금액2"]]  # fmt: skip
    if bank is not None:
        bank_sorted = sorted(bank, key=lambda r: (r["date"], r["time"]), reverse=True)
        transfers = [r for r in bank_sorted if r["out"] > 0]
        sheets.append(Sheet("계좌거래내역", bank_headers, [bank_row(r) for r in bank_sorted], (4, 5, 6)))
        sheets.append(Sheet("이체(출금)내역", bank_headers, [bank_row(r) for r in transfers], (4, 5, 6)))
        summary.append(["계좌 입금", sum(1 for r in bank if r["in"] > 0), sum(r["in"] for r in bank), ""])
        summary.append(["계좌 출금(이체)", len(transfers), sum(r["out"] for r in transfers), ""])
    if invoices is not None:
        for key, title, who in (("sales", "세금계산서_발급(매출)", "공급받는자"),
                                ("purchases", "세금계산서_수취(매입)", "공급자")):  # fmt: skip
            items = invoices.get(key, [])
            sheets.append(Sheet(title, invoice_headers(who), invoice_rows(items), (6, 7, 8)))
            summary.append([title.replace("세금계산서_", "세금계산서 "), len(items),
                            sum(i["supply_amount"] for i in items), sum(i["tax_amount"] for i in items)])  # fmt: skip
    summary.append(["", "", "", ""])
    summary.append(["※ 금액1 = 금액(세금계산서는 공급가액), 금액2 = 세액", "", "", ""])
    for message in warnings:
        summary.append(["⚠ " + message, "", "", ""])

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    for row in summary:
        ws.append(row)
    ws["A1"].font = Font(bold=True)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"
    ws.column_dimensions["A"].width = 28
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in sheets:
        ws = wb.create_sheet(sheet.title)
        ws.append(sheet.headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.rows:
            ws.append(row)
        if sheet.rows:
            total = ["합계"] + [""] * (len(sheet.headers) - 1)
            last = len(sheet.rows) + 1
            for col in sheet.money_cols:
                letter = get_column_letter(col)
                if not sheet.headers[col - 1].endswith("잔액"):
                    total[col - 1] = f"=SUM({letter}2:{letter}{last})"
            ws.append(total)
            for cell in ws[last + 1]:
                cell.font = Font(bold=True)
        for col in sheet.money_cols:
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "#,##0"
        for index, header in enumerate(sheet.headers, start=1):
            longest = max([len(str(header)) * 2] + [len(str(r[index - 1])) for r in sheet.rows[:500]])
            ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 40)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ 데모 데이터


def demo_data(start: date, end: date):
    day = lambda n: (end - timedelta(days=n)).strftime("%Y%m%d")  # noqa: E731
    iso = lambda n: (end - timedelta(days=n)).isoformat()  # noqa: E731
    bank = [
        {"account": "140000000000", "account_name": "기업자유예금", "date": day(0), "time": "101530",
         "out": 0, "in": 3300000, "balance": 15300000, "desc1": "(주)예시고객", "desc2": "타행이체", "desc3": "", "desc4": ""},
        {"account": "140000000000", "account_name": "기업자유예금", "date": day(2), "time": "143005",
         "out": 1100000, "in": 0, "balance": 12000000, "desc1": "예시상사", "desc2": "인터넷", "desc3": "9월자재대", "desc4": ""},
        {"account": "140000000000", "account_name": "기업자유예금", "date": day(5), "time": "090000",
         "out": 2500000, "in": 0, "balance": 13100000, "desc1": "급여", "desc2": "대량이체", "desc3": "", "desc4": ""},
    ]  # fmt: skip
    invoices = {
        "sales": [{"approval_number": "2026100100000000ABCDEF01", "written_date": iso(1), "issued_date": iso(1),
                   "transmitted_date": iso(0), "counterparty_name": "(주)예시고객", "item_name": "설계용역",
                   "supply_amount": 3000000, "tax_amount": 300000, "total_amount": 3300000}],
        "purchases": [{"approval_number": "2026100100000000ABCDEF02", "written_date": iso(3), "issued_date": iso(3),
                       "transmitted_date": iso(2), "counterparty_name": "예시상사", "item_name": "자재",
                       "supply_amount": 1000000, "tax_amount": 100000, "total_amount": 1100000}],
    }  # fmt: skip
    return bank, invoices


# ------------------------------------------------------------------ 진입점


def default_output(start: date, end: date) -> Path:
    out_dir = Path(env("OUTPUT_DIR") or data_dir() / "output").expanduser()
    return out_dir / f"거래_세금계산서_{start:%Y%m%d}-{end:%Y%m%d}.xlsx"


def fetch_bank_browser(start: date, end: date, log, done=None) -> tuple[list[dict], list[str]]:
    """CODEF 없이: 크롬 화면 반자동으로 엑셀을 받아 읽는다."""
    import bank_excel
    import shinhan_browser

    try:
        files = shinhan_browser.download_statements(start, end, data_dir(), log=log, done=done)
    except shinhan_browser.StepFailed as error:
        raise AppError(str(error)) from None
    return bank_excel.parse_files(files, start, end)


def fetch_bank_files(start: date, end: date, files) -> tuple[list[dict], list[str]]:
    import bank_excel

    if not files:
        raise AppError("은행 거래내역 파일을 고르세요.")
    return bank_excel.parse_files(files, start, end)


def run(start: date, end: date, *, cert_dir: Path | None = None, password: str = "",
        bank_on: bool = True, hometax_on: bool = True, demo: bool = False,
        out: Path | None = None, log=print, bank_mode: str = "browser",
        bank_files=None, bank_done=None) -> tuple[Path, list[str]]:
    """조회 → 엑셀 저장. (저장 경로, 경고 목록)을 돌려준다. CLI와 맥 앱이 함께 쓴다.

    bank_mode: browser(크롬 화면 반자동, 기본) / codef / files(내려받아 둔 엑셀 파일)"""
    out = out or default_output(start, end)
    log(f"조회 기간: {start} ~ {end}")
    warnings: list[str] = []
    bank = invoices = None
    if demo:
        bank, invoices = demo_data(start, end)
        warnings.append("데모 데이터입니다(실제 조회 아님).")
    else:
        needs_cert = hometax_on or (bank_on and bank_mode == "codef")
        if needs_cert and not cert_dir:
            raise AppError("공동인증서 폴더를 지정하세요.")
        if needs_cert and not password:
            raise AppError("공동인증서 비밀번호를 입력하세요.")
        if bank_on:
            log("[1/2] 신한은행 기업 계좌 거래내역 조회 중...")
            try:
                if bank_mode == "codef":
                    bank, bank_warnings = fetch_bank(start, end, cert_dir, password)
                elif bank_mode == "files":
                    bank, bank_warnings = fetch_bank_files(start, end, bank_files)
                else:
                    bank, bank_warnings = fetch_bank_browser(start, end, log, bank_done)
                warnings += bank_warnings
                log(f"  거래 {len(bank)}건")
            except Exception as error:  # 네트워크 오류 등도 한쪽 실패로 처리
                warnings.append(f"신한은행 조회 실패: {error}")
                log(f"  실패: {error}")
        if hometax_on:
            log("[2/2] 홈택스 세금계산서 발급·수취 내역 조회 중...")
            try:
                invoices = fetch_hometax(start, end, cert_dir, password)
                log(f"  발급 {len(invoices['sales'])}건, 수취 {len(invoices['purchases'])}건")
            except Exception as error:
                warnings.append(f"홈택스 조회 실패: {error}")
                log(f"  실패: {error}")
        if bank is None and invoices is None:
            raise AppError("조회에 모두 실패했습니다.\n" + "\n".join(warnings))
    build_workbook(out, start, end, bank, invoices, warnings)
    log(f"엑셀 저장: {out}")
    return out, warnings


def main(argv=None) -> int:
    load_dotenv(HERE / ".env")
    parser = argparse.ArgumentParser(description="신한은행 기업 거래·이체 + 홈택스 세금계산서 → 엑셀")
    parser.add_argument("--days", type=int, default=7, help="오늘 포함 최근 N일 (기본 7)")
    parser.add_argument("--from", dest="start", help="시작일 YYYY-MM-DD")
    parser.add_argument("--to", dest="end", help="종료일 YYYY-MM-DD")
    parser.add_argument("--skip-bank", action="store_true", help="은행 조회 생략")
    parser.add_argument("--skip-hometax", action="store_true", help="홈택스 조회 생략")
    parser.add_argument("--bank", choices=["browser", "codef", "files"], default=env("BANK_MODE", "browser"),
                        help="은행 조회 방식: browser(크롬 반자동, 기본) / codef / files")
    parser.add_argument("--bank-file", action="append", default=[], help="--bank files 일 때 거래내역 파일")
    parser.add_argument("--demo", action="store_true", help="가짜 데이터로 엑셀만 만들어 보기")
    parser.add_argument("--out", help="저장할 엑셀 파일 경로")
    args = parser.parse_args(argv)

    try:
        start, end = parse_period(args)
        cert_dir = password = None
        needs_cert = not args.skip_hometax or (not args.skip_bank and args.bank == "codef")
        if not args.demo and needs_cert:
            if not env("CERT_DIR"):
                raise AppError("CERT_DIR(공동인증서 폴더)를 설정하세요.")
            cert_dir = Path(env("CERT_DIR")).expanduser()
            password = env("CERT_PASSWORD") or getpass.getpass("공동인증서 비밀번호: ")
        _, warnings = run(
            start, end, cert_dir=cert_dir, password=password or "",
            bank_on=not args.skip_bank, hometax_on=not args.skip_hometax, demo=args.demo,
            out=Path(args.out) if args.out else None,
            bank_mode=args.bank, bank_files=args.bank_file,
        )  # fmt: skip
        return 1 if warnings and not args.demo else 0
    except AppError as error:
        print(f"오류: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
